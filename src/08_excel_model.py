"""
Aşama 8 — 1. Formüllerle çalışan Excel finansal modeli (FAST'a yakın yapı) ve Python mutabakatı.

Python'un verdiği (dispatch'e bağlı) yıllık birim büyüklükler "Python_girdileri" sayfasında; geri kalan her şey
(YEKDEM yıllık katılım kararı, gelir, kesinti, OPEX, amortisman, 5 yıllık FIFO zarar mahsubu, KDV, borç, DSCR/LLCR,
nakit akışı, NPV/IRR/LCOE) Excel formülleriyle hesaplanır. Senaryo seçimi Girdiler sayfasından.

Mutabakat: her senaryo (4 gelir yapısı × 3 fiyat × P50/P90 × teşvik var/yok) için Python sonucu doğrudan
src/risk/degerleme.nakit_akislari ile (aynı deterministik girdilerle) hesaplanır; Excel dosyası her senaryo için
LibreOffice'te yeniden hesaplatılır ve değerler karşılaştırılır. Sonuç "Senaryolar" ve "Kontroller" sayfalarına yazılır.

Çıktı: outputs/08/Ezine_Hibrit_Finansal_Model.xlsx, outputs/08/excel_mutabakat.csv
Çalıştırma: .venv/bin/python src/08_excel_model.py   (LibreOffice: ~/Applications/LibreOffice.app)
"""

import importlib.util
import itertools
import json
import shutil
import subprocess
import sys
import tempfile
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

from src.ayarlar import CIKTILAR, deger
from src.risk.degerleme import AI, irr_vektor, nakit_akislari
from src.risk.doyum import deterministik_yol, kurulu_yol
from src.risk.senaryo import h_yolu
from src.risk.tasarim import Tasarim
from src.risk.yekdem_tablo import yekdem_tablosu

CIKTI = CIKTILAR / "08"
DOSYA = CIKTI / "Ezine_Hibrit_Finansal_Model.xlsx"
SOFFICE = Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"
COD_BAZ = 2030
YIL0, NYIL = 2026, 39  # zaman çizelgesi: t = 1 → 2026 ... t = 39 → 2064; t = 0 bugün (edinim primi)
ALANLAR = ["gelir_res_q", "E_res", "gelir_d_q", "E_d", "yh_a", "rezerv", "E_k", "E_A"]
YAPILAR = ["Tüccar", "PPA", "YEKDEM nominal", "YEKDEM endeksli"]
SEVIYE = {"Düşük": 45.0, "Baz": 58.0, "Yüksek": 72.0}
MAVI, SIYAH = Font(color="1F4E9E"), Font(color="000000")
BASLIK = Font(bold=True, color="FFFFFF")
BASLIK_DOLGU = PatternFill("solid", fgColor="1D4F91")
ARA = PatternFill("solid", fgColor="EEF3FA")
GIRDI_DOLGU = PatternFill("solid", fgColor="FFF8DC")


def yukle(ad):
    spec = importlib.util.spec_from_file_location(ad, Path(__file__).parent / f"{ad}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


# =====================================================================================================================
# 1. Python tarafı
# =====================================================================================================================
class PythonModel:
    def __init__(self):
        st = yukle("05_stres_testi")
        self.st, o = st, st.gy.Ortam()
        self.o, cfg = o, o.cfg
        self.cfg, self.v = cfg, o.v
        C = Tasarim(**json.loads((CIKTILAR / "05_risk" / "adaylar.json").read_text())["C: gelir yığınlı optimal"])
        self.G = replace(C, ad="Nihai tasarım", bess_mw=30, bess_mwh=60)
        self.GY = replace(self.G, ppa=True)
        self.tb_T = st.doyum_tablosu(o, self.G)
        self.tb_Y = yekdem_tablosu(o, self.GY)
        N = o.v.omur
        self.N = N
        self.yillar = np.arange(COD_BAZ, COD_BAZ + N)
        fo = float(pd.read_csv(CIKTILAR / "05_risk" / "s_firsat_maliyeti.csv", index_col=0).loc["Fırsat maliyeti / 2026 seviyesi", "deger"])
        # deterministik baz: baz doyum yolu, ortak sürücü (baz kurulum yolu), beklenen kesinti
        os_ = deger(cfg["ortak_surucu"])
        kur = kurulu_yol(deger(cfg["stres_testi"]["senaryolar"])["baz"]["kurulu_mw"], self.yillar)
        emme = lambda mw, y: np.clip(
            mw
            * os_["sure_saat"]
            * os_["fazla_gun"]
            * os_["verim"]
            / (os_["fazla_2026_twh"] * 1e6 * (1 + os_["fazla_buyume"]) ** (y - 2026)),
            0,
            1,
        )
        a_abs = emme(kur, self.yillar)
        a26 = float(emme(1000.0, 2026))
        h0 = h_yolu(cfg, self.yillar, o.kan_taban, o.oran0)
        self.h = h0 + (o.oran0 - h0) * a_abs * os_["kanibalizasyon_geri_alma"]
        self.a_sfk = deterministik_yol(cfg, "baz", self.yillar, fo)
        k = deger(cfg["kesinti"])
        self.kappa = np.clip(np.mean(k["oran_ucgen"]) * (1 + k["gunes_buyume"]) ** (self.yillar - 2026) * (1 - a_abs) / (1 - a26), 0, 0.3)
        self.tazmin = float(np.mean(k["tazmin_orani"]))
        kesl = pd.read_csv(CIKTILAR / "07" / "kesinti_batarya_kurtarma.csv", index_col=0)
        self.kurtarma_m = float(kesl.loc["30 MW / 52,5 MWh", "Batarya kurtarma payı [%]"]) / 100
        self.deger_faktoru = float(
            kesl.loc["Bataryasız", "Kayıp [M USD]"] / (k["oran_2026"] * kesl.loc["Bataryasız", "Kesintisiz gelir [M USD]"])
        )
        self.kurtarma_y = 0.35
        self.m_p90 = 0.94
        yk = deger(cfg["yekdem"])
        self.yk_nom = float(np.mean([yk["taban_usd_sent"], yk["tepe_usd_sent"], yk["tavan_usd_sent"]]))  # ¢/kWh, üçgen ortalaması
        self.yk_sure = yk["sure_yil"]
        self.esik = deger(cfg["yekdem_opsiyon"])["esik_carpani"]
        self.pi = deger(cfg["finans"]["abd_enflasyonu"])
        self.birim = self._birimler()
        b = deger(cfg["borc"])
        self.borc_cfg = {k_: deger(b[k_]) for k_ in b}
        self.ppa = {k_: deger(v_) for k_, v_ in cfg["gelir_yigini"]["ppa"].items()}
        self.ts = deger(cfg["tesvik"])

    def _girdi(self, m, L=58.0):
        N = self.N
        return {
            "h": self.h[None],
            "m": np.full((1, N), m),
            "a": self.a_sfk[None],
            "L": np.full((1, N), L),
            "tavan": np.zeros((1, N), bool),
            "ws_ortalama": True,
            **{k: np.ones(1) for k in ["capex_res", "capex_ges", "capex_bess", "opex", "bozunma"]},
        }

    def _birimler(self):
        """Rejim × üretim durumu → (N, alan) birim büyüklükler (Python_girdileri)."""
        out = {}
        for durum, m in [("P50", 1.0), ("P90", self.m_p90)]:
            g = self._girdi(m)
            out[("T", durum)] = nakit_akislari(self.G, self.tb_T, self.o.c, g)["X"][0]
            for rej, ip in [("Y0", 0), ("Y1", 1)]:
                g2 = {**g, "ip": np.full((1, self.N), ip)}
                out[(rej, durum)] = nakit_akislari(self.GY, self.tb_Y, self.o.c, g2)["X"][0]
        return out

    def yekdem_reel(self, endeksli):
        return np.full(self.N, self.yk_nom * 10) if endeksli else self.yk_nom * 10 / (1 + self.pi) ** (self.yillar - 2026)

    def sonuc(self, yapi, seviye, durum, tesvik, prim=0.0):
        L = SEVIYE[seviye]
        N = self.N
        k = np.arange(1, N + 1)
        yek = yapi.startswith("YEKDEM")
        P = self.yekdem_reel(yapi == "YEKDEM endeksli") if yek else np.full(N, self.ppa["iskonto"] * 0 + 58.0 * (1 - self.ppa["iskonto"]))
        katil = (k <= self.yk_sure) & (P > self.esik * L) if yek else np.zeros(N, bool)

        def kos(d):
            if yek:
                X = np.where(katil[:, None], self.birim[("Y1", d)], self.birim[("Y0", d)])
            else:
                X = self.birim[("T", d)]
            phi = (
                katil.astype(float)
                if yek
                else (np.where(k <= self.ppa["vade"], self.ppa["pay"] * 0.94, 0.0) if yapi == "PPA" else np.zeros(N))
            )
            kg = np.where(
                katil,
                self.kappa * (1 - self.tazmin * self.deger_faktoru) * (1 - self.kurtarma_y),
                self.kappa * self.deger_faktoru * (1 - self.tazmin) * (1 - self.kurtarma_m),
            )
            ke = self.kappa * (1 - self.kurtarma_y)
            g = {
                **self._girdi(1.0 if d == "P50" else self.m_p90, L),
                "X": X[None],
                "phi": phi,
                "ppa_fiyat": P[None],
                "kesinti_gelir": kg[None],
                "kesinti_enerji": ke[None],
            }
            g.pop("ws_ortalama")
            if yek:
                g["desarj_sozlesmeli"] = True
            t = self.GY if yek else (replace(self.G, ppa=True) if yapi == "PPA" else self.G)
            return nakit_akislari(t, self.tb_T, self.o.c, g), phi

        r, phi = kos(durum)
        r90, _ = kos("P90")
        v = self.v
        nb = v.insaat_suresi
        capex = float(r["capex_top"][0])
        # KDV (teşvik belgesi yoksa)
        gelir_kdv = (r["gelir_spot"] + r["gelir_bess"] + np.atleast_2d(r["gelir_yh"]))[0]
        odenen = np.zeros(nb)
        mahsup = np.zeros(N)
        if not tesvik:
            kdv = self.ts["kdv_orani"] * self.ts["kdv_kapsami"] * capex
            bak = 0.0
            for j, pay in enumerate(v.capex_takvimi):
                odenen[j] = kdv * pay
                bak = bak * (1 - self.ts["kdv_reel_asinma"]) + odenen[j]
            for y in range(N):
                bak *= 1 - self.ts["kdv_reel_asinma"]
                mahsup[y] = min(bak, self.ts["kdv_orani"] * gelir_kdv[y])
                bak -= mahsup[y]
        # Proje nakit akışı (t = 0 prim; t = 1 → 2026; inşaat COD−2, COD−1)
        T = NYIL + 1
        cf = np.zeros(T)
        cf[0] = -prim
        bas = COD_BAZ - nb - YIL0 + 1
        cf[bas : bas + nb] = -capex * np.array(v.capex_takvimi) - odenen
        cf[bas + nb : bas + nb + N] = r["cfads"][0] + mahsup
        df = (1 + v.wacc) ** -np.arange(T)
        npv = float(cf @ df)
        irr = float(irr_vektor(cf[None])[0])
        # Borç (heykelsi; P90 CFADS; sözleşmeli pay DSCR harmanı)
        sozl = phi * P * (r["E_res"][0] + r["E_ges"][0])
        pay = np.clip(sozl / np.maximum(r["favok"][0], 1), 0, 1)
        dscr_h = 1.0 / (pay / self.ppa["dscr"] + (1 - pay) / self.borc_cfg["hedef_dscr"])
        vade = self.ppa["borc_vadesi"] if yapi == "PPA" else self.borc_cfg["vade"]
        i = self.borc_cfg["faiz_nominal"]
        enf = (1 + self.pi) ** (nb + k)
        servis = np.where(k <= vade, np.clip(r90["cfads"][0], 0, None) * enf / dscr_h, 0.0)
        borc_nom = float((servis * (1 + i) ** -k).sum())
        capex_nom = capex * float((np.array(v.capex_takvimi) * (1 + self.pi) ** (np.arange(nb) + 1)).sum())
        if borc_nom > self.borc_cfg["azami_borc_orani"] * capex_nom:
            servis *= self.borc_cfg["azami_borc_orani"] * capex_nom / borc_nom
            borc_nom = self.borc_cfg["azami_borc_orani"] * capex_nom
        bak, faiz = borc_nom, np.zeros(N)
        for y in range(N):
            faiz[y] = bak * i
            bak = bak + faiz[y] - servis[y]
        borc_reel = borc_nom / (1 + self.pi) ** nb
        kalkan = np.minimum(faiz * v.vergi, np.clip(r["vergi"][0], 0, None) * enf)
        oz = np.zeros(T)
        oz[0] = -prim
        oz[bas : bas + nb] = (-capex + borc_reel) * np.array(v.capex_takvimi) - odenen
        oz[bas + nb : bas + nb + N] = r["cfads"][0] - (servis - kalkan) / enf + mahsup
        oz_irr = float(irr_vektor(oz[None])[0])
        akt = servis > 0
        dscr = np.where(akt, r["cfads"][0] * enf / np.where(akt, servis, 1), np.nan)
        # LCOE (vergi öncesi, satılan enerji; inşaat başından iskonto)
        gelir_top = (r["gelir_spot"] + r["gelir_bess"] + np.atleast_2d(r["gelir_yh"]) + r["gelir_yekg"])[0]
        maliyet = gelir_top - r["vergi"][0] - r["cfads"][0]
        dfl = (1 + v.wacc) ** -(np.arange(nb + N) + 1.0)
        lcoe = float(((capex * np.array(v.capex_takvimi)) @ dfl[:nb] + maliyet @ dfl[nb:]) / (r["satilan"][0] @ dfl[nb:]))
        return {
            "NPV [M USD]": npv / 1e6,
            "Proje IRR [%]": 100 * irr,
            "Öz kaynak IRR [%]": 100 * oz_irr,
            "Asgari DSCR": float(np.nanmin(dscr)),
            "LCOE [USD/MWh]": lcoe,
            "Borç [M USD]": borc_reel / 1e6,
            "YEKDEM katılım yılı": int(katil.sum()),
        }


# =====================================================================================================================
# 2. Çalışma kitabı
# =====================================================================================================================
class Kitap:
    def __init__(self, pm: PythonModel):
        self.pm = pm
        self.wb = openpyxl.Workbook()
        self.ad = {}  # isim → mutlak hücre referansı
        self.satir = {}  # (sayfa, anahtar) → satır
        self.C0 = 5  # E sütunu: t = 0
        self.son = self.C0 + NYIL

    # --- yardımcılar
    def ws(self, ad):
        if ad in self.wb.sheetnames:
            return self.wb[ad]
        w = self.wb.create_sheet(ad)
        w.sheet_view.showGridLines = False
        return w

    def kol(self, j):  # j = 0 … NYIL (t)
        return CL(self.C0 + j)

    def r(self, sayfa, anahtar, j):  # zaman serisi hücresi
        return f"'{sayfa}'!{self.kol(j)}{self.satir[(sayfa, anahtar)]}"

    def aralik(self, sayfa, anahtar, j0=1, j1=NYIL):
        rw = self.satir[(sayfa, anahtar)]
        return f"'{sayfa}'!${self.kol(j0)}${rw}:${self.kol(j1)}${rw}"

    def baslik_satiri(self, w, satir, metin):
        w.cell(satir, 1, metin).font = Font(bold=True, size=11, color="1D4F91")

    def zaman_basligi(self, w, satir=4):
        w.cell(satir, 1, "Kalem").font = BASLIK
        w.cell(satir, 2, "Birim").font = BASLIK
        w.cell(satir, 3, "Not / kaynak").font = BASLIK
        w.cell(satir, 4, "Toplam / değer").font = BASLIK
        for j in range(NYIL + 1):
            c = w.cell(satir, self.C0 + j, "t0 (bugün)" if j == 0 else YIL0 + j - 1)
            c.font = BASLIK
        for cc in range(1, self.son + 1):
            w.cell(satir, cc).fill = BASLIK_DOLGU
        w.column_dimensions["A"].width = 46
        w.column_dimensions["B"].width = 12
        w.column_dimensions["C"].width = 40
        w.column_dimensions["D"].width = 14
        for j in range(NYIL + 1):
            w.column_dimensions[self.kol(j)].width = 11
        w.freeze_panes = w.cell(satir + 1, self.C0)

    def seri(self, sayfa, anahtar, etiket, birim, formul, not_="", toplam=True, fmt="#,##0", j0=0, deger_t0=0):
        """formul(j) → hücre formülü (j = 1…NYIL); t0 sütunu deger_t0."""
        w = self.ws(sayfa)
        rw = w.max_row + 1
        assert self.satir.get((sayfa, anahtar), rw) == rw, (sayfa, anahtar)
        self.satir[(sayfa, anahtar)] = rw
        w.cell(rw, 1, etiket)
        w.cell(rw, 2, birim)
        w.cell(rw, 3, not_)
        for j in range(j0, NYIL + 1):
            c = w.cell(rw, self.C0 + j, deger_t0 if j == 0 else formul(j))
            c.number_format = fmt
        if toplam:
            w.cell(rw, 4, f"=SUM({self.kol(1)}{rw}:{self.kol(NYIL)}{rw})").number_format = fmt
        return rw

    def girdi(self, w, rw, ad, etiket, deger_, birim, kaynak, fmt="#,##0.00", liste=None):
        w.cell(rw, 1, etiket)
        c = w.cell(rw, 3, deger_)
        c.font = MAVI
        c.fill = GIRDI_DOLGU
        c.number_format = fmt
        w.cell(rw, 4, birim)
        w.cell(rw, 5, kaynak)
        self.ad[ad] = f"Girdiler!$C${rw}"
        if liste:
            dv = DataValidation(type="list", formula1='"' + ",".join(liste) + '"', allow_blank=False)
            w.add_data_validation(dv)
            dv.add(c)
        return rw + 1

    # --- sayfalar
    def kapak(self):
        w = self.wb.active
        w.title = "Kapak"
        w.sheet_view.showGridLines = False
        w["A1"] = "Ezine Hibrit RES + Batarya — Finansal Model"
        w["A1"].font = Font(bold=True, size=16, color="1D4F91")
        w["A2"] = "Reel 2026 USD · %7 reel WACC · FAST standardına yakın yapı · Ekim 2026"
        w["A4"] = "Renk kodu"
        w["A4"].font = Font(bold=True)
        w["A5"] = "Mavi yazı, sarı zemin = girdi"
        w["A5"].font = MAVI
        w["A5"].fill = GIRDI_DOLGU
        w["A6"] = "Siyah yazı = formül"
        w["A7"] = "Python_girdileri = Python modelinden (saatlik dispatch) gelen değerler"
        w["A9"] = "İçindekiler"
        w["A9"].font = Font(bold=True)
        icerik = [
            ("Girdiler", "Tüm varsayımlar ve senaryo seçimi (gelir yapısı, fiyat, P50/P90, teşvik, edinim primi)"),
            ("Python_girdileri", "Yıllık üretim ve batarya büyüklükleri (Python dispatch LP'leri), kaynak açıklamalı"),
            ("Zaman", "Zaman çizelgesi (COD girdiden), iskonto, YEKDEM yıllık katılım kararı"),
            ("Uretim", "Seçilen yapı ve üretim durumuna göre yıllık üretim, sistem kesintisi"),
            ("Gelir", "RES ve batarya gelir kalemleri (YEKDEM dönemi ve sonrası)"),
            ("OPEX", "İşletme giderleri, FAVÖK, augmentation/yenileme, söküm"),
            ("Vergi", "Amortisman, 5 yıllık FIFO zarar mahsubu, kurumlar vergisi, KDV (teşvik yoksa)"),
            ("Borc", "DSCR hedefli heykelsi borç, geri ödeme tablosu, DSCR, LLCR"),
            ("Nakit_akisi", "Proje ve öz kaynak nakit akışları"),
            ("Ozet", "NPV, IRR, LCOE, DSCR, geri ödeme, azami edinim primi"),
            ("Senaryolar", "48 senaryo: Python ve Excel (LibreOffice ile yeniden hesaplanmış) sonuçları"),
            ("Kontroller", "Nakit tutarlılığı, borç sıfırlanması, Python mutabakatı, formül hatası sayımı"),
        ]
        for i, (s, a) in enumerate(icerik):
            c = w.cell(10 + i, 1, s)
            c.hyperlink = f"#'{s}'!A1"
            c.font = Font(color="1F4E9E", underline="single")
            w.cell(10 + i, 2, a)
        w["A24"] = (
            "Not: Monte Carlo sonuçları (beklenen NPV, P(NPV>0)) bu deterministik modelde yoktur; Ozet sayfasında Python "
            "Monte Carlo referansı ayrıca verilir."
        )
        w.column_dimensions["A"].width = 30
        w.column_dimensions["B"].width = 110

    def girdiler(self):
        pm, v = self.pm, self.pm.v
        w = self.ws("Girdiler")
        w["A1"] = "Girdiler"
        w["A1"].font = Font(bold=True, size=14, color="1D4F91")
        w["A2"] = "Mavi = girdi. Senaryo seçimi bu sayfadan; değerler config.yaml'dan (kaynaklarıyla)."
        for i, h in enumerate(["Girdi", "", "Değer", "Birim", "Kaynak / not"]):
            c = w.cell(4, i + 1, h)
            c.font = BASLIK
            c.fill = BASLIK_DOLGU
        w.column_dimensions["A"].width = 44
        w.column_dimensions["C"].width = 16
        w.column_dimensions["D"].width = 16
        w.column_dimensions["E"].width = 90
        rw = 5
        self.baslik_satiri(w, rw, "Senaryo seçimi")
        rw += 1
        rw = self.girdi(
            w,
            rw,
            "yapi",
            "Gelir yapısı",
            "YEKDEM nominal",
            "liste",
            "Baz vaka: YEKDEM yıllık opsiyonu (nominal koridor) + tüccar",
            "@",
            YAPILAR,
        )
        rw = self.girdi(
            w, rw, "seviye", "Fiyat seviyesi", "Baz", "liste", "Düşük 45 / Baz 58 / Yüksek 72 USD/MWh (reel)", "@", list(SEVIYE)
        )
        rw = self.girdi(w, rw, "durum", "Üretim durumu", "P50", "liste", "P90 = 10 yıllık P90/P50 = 0,94 (Aşama 5)", "@", ["P50", "P90"])
        rw = self.girdi(
            w,
            rw,
            "tesvik",
            "Yatırım teşvik belgesi",
            "Var",
            "liste",
            "CK 9903: KDV istisnası; yoksa KDV ön finansmanı",
            "@",
            ["Var", "Yok"],
        )
        rw = self.girdi(w, rw, "prim", "Edinim primi (bugünkü değer)", 0.0, "USD", "t = 0'da ödenir; azami prim = primsiz NPV", "#,##0")
        rw = self.girdi(
            w,
            rw,
            "cod",
            "COD yılı (ticari işletme başlangıcı)",
            COD_BAZ,
            "yıl",
            "Takvim P50 Kasım 2029 → yıllık modelde 2030 (Python_girdileri bu COD için)",
            "0",
        )
        rw += 1
        self.baslik_satiri(w, rw, "Fiyat seviyeleri")
        rw += 1
        for ad_, d in SEVIYE.items():
            rw = self.girdi(w, rw, f"L_{ad_}", f"Fiyat seviyesi — {ad_}", d, "USD/MWh", "config: fiyat_senaryolari.seviye")
        w.cell(rw, 1, "Seçilen fiyat seviyesi L")
        w.cell(
            rw,
            3,
            f'=IF({self.ad["seviye"]}="Düşük",{self.ad["L_Düşük"]},IF({self.ad["seviye"]}="Yüksek",{self.ad["L_Yüksek"]},{self.ad["L_Baz"]}))',
        )
        w.cell(rw, 4, "USD/MWh")
        self.ad["L"] = f"Girdiler!$C${rw}"
        rw += 1
        w.cell(rw, 1, "Yapı indeksi (1 Tüccar, 2 PPA, 3 YEKDEM nom., 4 YEKDEM end.)")
        w.cell(rw, 3, "=" + "+".join(f'({self.ad["yapi"]}="{y}")*{i + 1}' for i, y in enumerate(YAPILAR)))
        self.ad["yapi_i"] = f"Girdiler!$C${rw}"
        rw += 1
        w.cell(rw, 1, "Senaryo anahtarı (mutabakat)")
        w.cell(rw, 3, f'={self.ad["yapi"]}&"|"&{self.ad["seviye"]}&"|"&{self.ad["durum"]}&"|"&{self.ad["tesvik"]}')
        self.ad["anahtar"] = f"Girdiler!$C${rw}"
        rw += 2
        self.baslik_satiri(w, rw, "Tasarım")
        rw += 1
        t = pm.G
        rw = self.girdi(w, rw, "res_mw", "RES kurulu gücü (14 × 4,2 MW)", t.res_mw, "MW", "Aşama 4 optimizasyonu")
        rw = self.girdi(w, rw, "bess_mw", "Batarya gücü", t.bess_mw, "MW", "Ö1: tam 30 MW SFK rezervi için 2 saat")
        rw = self.girdi(w, rw, "bess_mwh", "Batarya enerjisi", t.bess_mwh, "MWh", "")
        rw = self.girdi(w, rw, "sebeke_mw", "Şebeke bağlantısı", t.sebeke_mw, "MW", "")
        rw += 1
        self.baslik_satiri(w, rw, "CAPEX ve teknik")
        rw += 1
        rw = self.girdi(w, rw, "capex_res", "RES CAPEX", v.capex_res, "USD/MW", "config: teknoloji.res.capex (IRENA/BNEF)", "#,##0")
        rw = self.girdi(
            w, rw, "capex_bess_e", "Batarya CAPEX — enerji", v.capex_bess_enerji, "USD/MWh", "config: teknoloji.bess.capex_enerji", "#,##0"
        )
        rw = self.girdi(
            w, rw, "capex_bess_p", "Batarya CAPEX — güç", v.capex_bess_guc, "USD/MW", "config: teknoloji.bess.capex_guc", "#,##0"
        )
        rw = self.girdi(w, rw, "capex_seb", "Şebeke CAPEX", v.capex_sebeke, "USD/MW", "config: teknoloji.sebeke.capex", "#,##0")
        rw = self.girdi(w, rw, "takvim1", "CAPEX takvimi — 1. inşaat yılı", v.capex_takvimi[0], "pay", "config: proje.capex_takvimi")
        rw = self.girdi(w, rw, "takvim2", "CAPEX takvimi — 2. inşaat yılı", v.capex_takvimi[1], "pay", "")
        rw = self.girdi(w, rw, "omur", "İşletme ömrü", v.omur, "yıl", "", "0")
        rw += 1
        self.baslik_satiri(w, rw, "İşletme giderleri")
        rw += 1
        rw = self.girdi(w, rw, "opex_res", "RES sabit O&M", v.opex_res, "USD/MW-yıl", "config", "#,##0")
        rw = self.girdi(w, rw, "arazi_res", "RES arazi", v.arazi_res, "USD/MW-yıl", "config", "#,##0")
        rw = self.girdi(w, rw, "opex_bess", "Batarya O&M (CAPEX'e oran)", v.opex_bess_oran, "oran", "config")
        rw = self.girdi(w, rw, "sigorta", "Sigorta (CAPEX'e oran)", v.sigorta_oran, "oran", "config", "0.0000")
        rw = self.girdi(w, rw, "teias_k", "TEİAŞ kapasite bedeli", v.teias_kapasite, "USD/MW-yıl", "config", "#,##0")
        rw = self.girdi(w, rw, "teias_e", "TEİAŞ enerji bedeli", v.teias_enerji, "USD/MWh", "config")
        rw = self.girdi(w, rw, "yipranma", "Batarya yıpranma maliyeti", v.bess_yipranma, "USD/MWh deşarj", "config")
        rw = self.girdi(
            w,
            rw,
            "deng",
            "RES dengesizlik maliyeti (gelire oran)",
            v.dengesizlik["res"],
            "oran",
            "config; YEKDEM'de de tesiste (YEK Yön. md. 18)",
        )
        rw = self.girdi(w, rw, "esz", "Eşzamanlılık düzeltmesi (RES, tüccar payı)", v.eszamanlilik_res, "oran", "config")
        rw = self.girdi(w, rw, "yekg", "YEK-G geliri", pm.o.c.yekg, "USD/MWh", "EPİAŞ YEK-G")
        rw = self.girdi(w, rw, "dz_azami", "Dengesizlik azaltımı — azami oran", pm.o.c.dz_azami, "oran", "TAHMİN")
        rw = self.girdi(w, rw, "dz_doyum", "Dengesizlik azaltımı — doyum (batarya/RES MW)", pm.o.c.dz_doyum, "oran", "TAHMİN")
        rw = self.girdi(w, rw, "aug_oran", "Augmentation (8, 16, 28. yıl; enerji CAPEX'ine oran)", v.augmentation["oran"], "oran", "config")
        rw = self.girdi(
            w, rw, "aug_dusus", "Batarya maliyet düşüşü (aug/yenileme)", v.augmentation["maliyet_dususu_yillik"], "oran/yıl", "config"
        )
        rw = self.girdi(w, rw, "omur_bess", "Batarya ömrü (yenileme yılı)", v.omur_bess, "yıl", "config", "0")
        rw = self.girdi(w, rw, "sokum_bess", "Söküm — batarya", v.sokum["bess"], "USD/MWh", "config", "#,##0")
        rw = self.girdi(w, rw, "sokum_res", "Söküm — RES (net)", v.sokum["res"], "USD/MW", "config", "#,##0")
        rw += 1
        self.baslik_satiri(w, rw, "Sistem kesintisi (TEİAŞ YAT)")
        rw += 1
        rw = self.girdi(w, rw, "tazmin", "Tazmin oranı (beklenen)", pm.tazmin, "oran", "TAHMİN, U(0,1) ortalaması (docs/07 §1)")
        rw = self.girdi(
            w,
            rw,
            "deger_f",
            "Kesilen enerjinin değer faktörü / PTF_talimat/YEKF",
            pm.deger_faktoru,
            "oran",
            "Talimat saatlerinde fiyat ortalamanın altında",
        )
        rw = self.girdi(w, rw, "kur_m", "Batarya kurtarma payı — tüccar (gelir)", pm.kurtarma_m, "oran", "Kesinti LP'si")
        rw = self.girdi(
            w, rw, "kur_y", "Batarya kurtarma payı — YEKDEM (enerji)", pm.kurtarma_y, "oran", "Kurtarılan enerji YEKDEM fiyatından ödenir"
        )
        rw += 1
        self.baslik_satiri(w, rw, "Gelir yapıları")
        rw += 1
        rw = self.girdi(w, rw, "ppa_pay", "PPA payı (%70 × P90/P50)", pm.ppa["pay"] * 0.94, "oran", "config: gelir_yigini.ppa")
        rw = self.girdi(w, rw, "ppa_vade", "PPA süresi", pm.ppa["vade"], "yıl", "", "0")
        rw = self.girdi(w, rw, "ppa_fiyat", "PPA fiyatı (58 × (1 − %10))", 58.0 * (1 - pm.ppa["iskonto"]), "USD/MWh", "TAHMİN iskonto")
        rw = self.girdi(
            w, rw, "yk_nom", "YEKDEM nominal fiyatı (koridor içi, beklenen)", pm.yk_nom, "USD ¢/kWh", "CK 7189: 4,95–6,05; üçgen ortalaması"
        )
        rw = self.girdi(w, rw, "yk_sure", "YEKDEM hak süresi (COD'dan)", pm.yk_sure, "yıl", "YEK Yön. md. 8/1", "0")
        rw = self.girdi(
            w,
            rw,
            "esik",
            "YEKDEM katılım eşiği (YEKDEM reel > eşik × beklenen PTF)",
            pm.esik,
            "kat",
            "config: yekdem_opsiyon (karar 30 Kasım t−1)",
            "0.000",
        )
        rw += 1
        self.baslik_satiri(w, rw, "Finansman ve vergi")
        rw += 1
        rw = self.girdi(w, rw, "wacc", "Reel WACC", v.wacc, "oran", "config", "0.0%")
        rw = self.girdi(w, rw, "pi", "ABD enflasyonu", pm.pi, "oran", "config", "0.0%")
        rw = self.girdi(w, rw, "vergi", "Kurumlar vergisi", v.vergi, "oran", "config", "0%")
        rw = self.girdi(w, rw, "amort", "Amortisman süresi", 10, "yıl", "config (tüm sınıflar 10 yıl)", "0")
        rw = self.girdi(w, rw, "dscr_t", "Hedef DSCR — tüccar payı", pm.borc_cfg["hedef_dscr"], "kat", "config: borc")
        rw = self.girdi(w, rw, "dscr_s", "Hedef DSCR — sözleşmeli pay (PPA/YEKDEM)", pm.ppa["dscr"], "kat", "config: gelir_yigini.ppa.dscr")
        rw = self.girdi(w, rw, "vade_t", "Borç vadesi (tüccar/YEKDEM)", pm.borc_cfg["vade"], "yıl", "config: borc.vade", "0")
        rw = self.girdi(w, rw, "vade_p", "Borç vadesi (PPA)", pm.ppa["borc_vadesi"], "yıl", "config: gelir_yigini.ppa.borc_vadesi", "0")
        rw = self.girdi(w, rw, "faiz", "Nominal faiz", pm.borc_cfg["faiz_nominal"], "oran", "config", "0.0%")
        rw = self.girdi(w, rw, "azami_borc", "Azami borç / CAPEX", pm.borc_cfg["azami_borc_orani"], "oran", "config", "0%")
        rw = self.girdi(w, rw, "kdv", "KDV oranı", pm.ts["kdv_orani"], "oran", "config: tesvik", "0%")
        rw = self.girdi(w, rw, "kdv_kapsam", "KDV'ye tabi CAPEX payı", pm.ts["kdv_kapsami"], "oran", "TAHMİN")
        rw = self.girdi(w, rw, "kdv_asinma", "Devreden KDV'nin reel aşınması", pm.ts["kdv_reel_asinma"], "oran/yıl", "TAHMİN (TL)")
        rw += 1
        self.baslik_satiri(w, rw, "Türetilmiş")
        rw += 1
        form = {
            "capex_res_t": ("RES CAPEX toplam", f"={self.ad['res_mw']}*{self.ad['capex_res']}"),
            "capex_bess_t": (
                "Batarya CAPEX toplam",
                f"={self.ad['bess_mw']}*{self.ad['capex_bess_p']}+{self.ad['bess_mwh']}*{self.ad['capex_bess_e']}",
            ),
            "capex_seb_t": ("Şebeke CAPEX toplam", f"={self.ad['sebeke_mw']}*{self.ad['capex_seb']}"),
        }
        for k_, (e, f) in form.items():
            w.cell(rw, 1, e)
            w.cell(rw, 3, f).number_format = "#,##0"
            w.cell(rw, 4, "USD")
            self.ad[k_] = f"Girdiler!$C${rw}"
            rw += 1
        w.cell(rw, 1, "CAPEX toplam")
        w.cell(rw, 3, f"={self.ad['capex_res_t']}+{self.ad['capex_bess_t']}+{self.ad['capex_seb_t']}").number_format = "#,##0"
        w.cell(rw, 4, "USD")
        self.ad["capex"] = f"Girdiler!$C${rw}"
        rw += 1
        w.cell(rw, 1, "Yapı YEKDEM mi?")
        w.cell(rw, 3, f"=IF({self.ad['yapi_i']}>=3,1,0)")
        self.ad["yek"] = f"Girdiler!$C${rw}"
        rw += 1
        w.cell(rw, 1, "Borç vadesi (seçilen yapı)")
        w.cell(rw, 3, f"=IF({self.ad['yapi_i']}=2,{self.ad['vade_p']},{self.ad['vade_t']})")
        self.ad["vade"] = f"Girdiler!$C${rw}"
        rw += 1
        w.cell(rw, 1, "Nominal CAPEX (borç tavanı için)")
        w.cell(
            rw, 3, f"={self.ad['capex']}*({self.ad['takvim1']}*(1+{self.ad['pi']})+{self.ad['takvim2']}*(1+{self.ad['pi']})^2)"
        ).number_format = "#,##0"
        self.ad["capex_nom"] = f"Girdiler!$C${rw}"
        rw += 1

    def python_girdileri(self):
        pm = self.pm
        w = self.ws("Python_girdileri")
        w["A1"] = "Python girdileri — saatlik dispatch LP'lerinden yıllık birim büyüklükler (işletme yılı 1–30, COD 2030)"
        w["A1"].font = Font(bold=True, size=13, color="1D4F91")
        w["A2"] = (
            "Kaynak: src/08_excel_model.py → src/risk/degerleme.nakit_akislari (X matrisi). Deterministik baz: 10 hava yılı × 2 fiyat "
            "şekli ortalaması, kanibalizasyon 'devam' + ortak batarya sürücüsü, baz doyum SFK yolu, beklenen sistem kesintisi. "
            "*_q = Σ üretim × birim fiyat şekli (fiyat seviyesi 1 için; gelir = L × *_q), E_* = MWh, yh_a = SFK birimi (gelir = a × yh_a)."
        )
        w["A3"] = (
            "Rejimler: T = tüccar dispatch'i (tüccar ve finansal PPA), Y0 = YEKDEM tablosunun katılımsız dilimi, Y1 = YEKDEM'e katılım "
            "(deşarj da YEKDEM fiyatından → arbitraj yok). P90: RES çarpanı 0,94. COD değişirse bu takvim bağımlı girdiler kaydırılmaz."
        )
        basl = ["İşletme yılı", "Takvim yılı (COD 2030)", "SFK fiyat çarpanı a", "Beklenen kesinti oranı κ"]
        kolon = {}
        c = 5
        for rej in ["T", "Y0", "Y1"]:
            for durum in ["P50", "P90"]:
                for f in ALANLAR:
                    basl.append(f"{rej}_{durum}_{f}")
                    kolon[(rej, durum, f)] = c
                    c += 1
        for i, h in enumerate(basl):
            x = w.cell(5, i + 1, h)
            x.font = BASLIK
            x.fill = BASLIK_DOLGU
            x.alignment = Alignment(wrap_text=True)
            w.column_dimensions[CL(i + 1)].width = 13
        for y in range(pm.N):
            rw = 6 + y
            w.cell(rw, 1, y + 1)
            w.cell(rw, 2, int(pm.yillar[y]))
            w.cell(rw, 3, float(pm.a_sfk[y])).number_format = "0.0000"
            w.cell(rw, 4, float(pm.kappa[y])).number_format = "0.0000"
            for (rej, durum, f), cc in kolon.items():
                x = w.cell(rw, cc, float(pm.birim[(rej, durum)][y, AI[f]]))
                x.number_format = "#,##0"
                x.font = MAVI
        self.pg_kolon = kolon
        self.pg_aralik = lambda cc: f"Python_girdileri!${CL(cc)}$6:${CL(cc)}${5 + pm.N}"
        self.ad["pg_a"] = self.pg_aralik(3)
        self.ad["pg_kappa"] = self.pg_aralik(4)

    def zaman(self):
        A = self.ad
        Z = "Zaman"
        w = self.ws(Z)
        w["A1"] = "Zaman çizelgesi ve YEKDEM yıllık katılım kararı"
        w["A1"].font = Font(bold=True, size=13, color="1D4F91")
        self.zaman_basligi(w)
        self.seri(Z, "yil", "Takvim yılı", "yıl", lambda j: YIL0 + j - 1, toplam=False, fmt="0", deger_t0=YIL0 - 1)
        self.seri(Z, "t", "İskonto dönemi t", "", lambda j: j, toplam=False, fmt="0", deger_t0=0)
        self.seri(
            Z,
            "insaat_pay",
            "İnşaat CAPEX payı",
            "pay",
            lambda j: f"=IF({self.r(Z, 'yil', j)}={A['cod']}-2,{A['takvim1']},IF({self.r(Z, 'yil', j)}={A['cod']}-1,{A['takvim2']},0))",
            fmt="0.00",
        )
        self.seri(
            Z,
            "k",
            "İşletme yılı k",
            "",
            lambda j: (
                f"=IF(AND({self.r(Z, 'yil', j)}>={A['cod']},{self.r(Z, 'yil', j)}<{A['cod']}+{A['omur']}),{self.r(Z, 'yil', j)}-{A['cod']}+1,0)"
            ),
            toplam=False,
            fmt="0",
        )
        self.seri(
            Z,
            "df",
            "İskonto faktörü (1 + WACC)^−t",
            "",
            lambda j: f"=1/(1+{A['wacc']})^{self.r(Z, 't', j)}",
            toplam=False,
            fmt="0.0000",
            deger_t0=1,
        )
        self.seri(
            Z,
            "enf",
            "Enflasyon endeksi (1 + π)^(2 + k)",
            "",
            lambda j: f"=IF({self.r(Z, 'k', j)}>0,(1+{A['pi']})^(2+{self.r(Z, 'k', j)}),0)",
            toplam=False,
            fmt="0.0000",
        )
        self.seri(
            Z,
            "dfl",
            "LCOE iskontosu (inşaat başından)",
            "",
            lambda j: f"=IF({self.r(Z, 'yil', j)}>={A['cod']}-2,1/(1+{A['wacc']})^({self.r(Z, 'yil', j)}-({A['cod']}-2)+1),0)",
            toplam=False,
            fmt="0.0000",
        )
        self.seri(
            Z,
            "pg_ok",
            "Python girdisi var mı (1 ≤ k ≤ 30)",
            "",
            lambda j: f"=IF(AND({self.r(Z, 'k', j)}>=1,{self.r(Z, 'k', j)}<=30),1,0)",
            toplam=False,
            fmt="0",
        )
        self.seri(
            Z,
            "yk_reel",
            "YEKDEM reel fiyatı",
            "USD/MWh",
            lambda j: f"=IF({A['yapi_i']}=4,{A['yk_nom']}*10,{A['yk_nom']}*10/(1+{A['pi']})^({self.r(Z, 'yil', j)}-2026))",
            toplam=False,
            fmt="0.00",
            not_="Nominal koridor reel olarak aşınır; endeksli varsayımda sabit",
        )
        self.seri(
            Z,
            "katil",
            "YEKDEM'e katılım (karar 30 Kasım t−1)",
            "0/1",
            lambda j: (
                f"=IF(AND({A['yek']}=1,{self.r(Z, 'k', j)}>=1,{self.r(Z, 'k', j)}<={A['yk_sure']},{self.r(Z, 'yk_reel', j)}>{A['esik']}*{A['L']}),1,0)"
            ),
            fmt="0",
            not_="Katıl ⇔ YEKDEM reel > eşik × beklenen PTF (deterministik: L)",
        )
        self.seri(
            Z,
            "phi",
            "Sözleşmeli pay φ",
            "",
            lambda j: (
                f"=IF({A['yapi_i']}=2,IF(AND({self.r(Z, 'k', j)}>=1,{self.r(Z, 'k', j)}<={A['ppa_vade']}),{A['ppa_pay']},0),{self.r(Z, 'katil', j)})"
            ),
            toplam=False,
            fmt="0.000",
        )
        self.seri(
            Z,
            "pf",
            "Sözleşmeli fiyat (PPA / YEKDEM)",
            "USD/MWh",
            lambda j: f"=IF({A['yapi_i']}=2,{A['ppa_fiyat']},{self.r(Z, 'yk_reel', j)})",
            toplam=False,
            fmt="0.00",
        )

    def _pg(self, rej_formul, durum_formul, f, j):
        """Python girdisi araması: rejim ve üretim durumuna göre INDEX(…, k)."""
        Z = "Zaman"

        def idx(rej, durum):
            return f"INDEX({self.pg_aralik(self.pg_kolon[(rej, durum, f)])},{self.r(Z, 'k', j)})"

        def durumlu(rej):
            if durum_formul == "P90":
                return idx(rej, "P90")
            return f'IF({self.ad["durum"]}="P90",{idx(rej, "P90")},{idx(rej, "P50")})'

        yek = f"IF({self.r(Z, 'katil', j)}=1,{durumlu('Y1')},{durumlu('Y0')})"
        return f"=IF({self.r(Z, 'pg_ok', j)}=1,IF({self.ad['yek']}=1,{yek},{durumlu('T')}),0)"

    def blok(self, b):
        """b = 'Sec' (seçilen üretim durumu) ya da 'P90' (borç boyutlandırma)."""
        A, Z = self.ad, "Zaman"
        U, G, O, V = "Uretim", "Gelir", "OPEX", "Vergi"
        etiket = "seçilen üretim durumu" if b == "Sec" else "P90 (borç boyutlandırma)"
        for s in (U, G, O, V):
            w = self.ws(s)
            if w.max_row < 4:
                w["A1"] = {
                    "Uretim": "Üretim ve sistem kesintisi",
                    "Gelir": "Gelir",
                    "OPEX": "İşletme giderleri ve FAVÖK",
                    "Vergi": "Amortisman, zarar mahsubu, vergi ve KDV",
                }[s]
                w["A1"].font = Font(bold=True, size=13, color="1D4F91")
                self.zaman_basligi(w)
            rw = w.max_row + 2
            w.cell(rw, 1, f"— Blok: {etiket} —").font = Font(bold=True, color="1D4F91")
            w.cell(rw, 1).fill = ARA
        df_ = "P90" if b == "P90" else None
        for f in ALANLAR:
            self.seri(
                U,
                f"{b}_{f}",
                f"{f}",
                "MWh" if f.startswith("E") or f == "rezerv" else "birim",
                lambda j, f=f: self._pg(None, df_, f, j),
                not_="Python_girdileri (rejim: katılıma göre Y1/Y0, aksi T)",
            )
        self.seri(
            U,
            f"{b}_kappa",
            "Sistem kesintisi oranı κ",
            "oran",
            lambda j: f"=IF({self.r(Z, 'pg_ok', j)}=1,INDEX({A['pg_kappa']},{self.r(Z, 'k', j)}),0)",
            toplam=False,
            fmt="0.000",
        )
        self.seri(
            U,
            f"{b}_kg",
            "Kesinti gelir kaybı payı",
            "oran",
            lambda j: (
                f"=IF({self.r(Z, 'katil', j)}=1,{self.r(U, f'{b}_kappa', j)}*(1-{A['tazmin']}*{A['deger_f']})*(1-{A['kur_y']}),"
                f"{self.r(U, f'{b}_kappa', j)}*{A['deger_f']}*(1-{A['tazmin']})*(1-{A['kur_m']}))"
            ),
            toplam=False,
            fmt="0.0000",
            not_="YEKDEM yılı: YEK farkı ödenmez; tüccar: DGP tazmini",
        )
        self.seri(
            U,
            f"{b}_ke",
            "Kesinti enerji kaybı payı",
            "oran",
            lambda j: f"={self.r(U, f'{b}_kappa', j)}*(1-{A['kur_y']})",
            toplam=False,
            fmt="0.0000",
        )
        self.seri(
            U,
            f"{b}_E_res_net",
            "RES satılan enerji (kesinti sonrası)",
            "MWh",
            lambda j: f"={self.r(U, f'{b}_E_res', j)}*(1-{self.r(U, f'{b}_ke', j)})",
        )
        self.seri(
            U, f"{b}_satilan", "Toplam satılan enerji", "MWh", lambda j: f"={self.r(U, f'{b}_E_res_net', j)}+{self.r(U, f'{b}_E_d', j)}"
        )
        # Gelir
        phi, pf = (lambda j: self.r(Z, "phi", j)), (lambda j: self.r(Z, "pf", j))
        self.seri(
            G,
            f"{b}_g_res",
            "RES geliri (tüccar + sözleşmeli, kesinti sonrası)",
            "USD",
            lambda j: (
                f"=((1-{phi(j)})*{A['L']}*{self.r(U, f'{b}_gelir_res_q', j)}+{phi(j)}*{pf(j)}*{self.r(U, f'{b}_E_res', j)})*(1-{self.r(U, f'{b}_kg', j)})"
            ),
        )
        self.seri(
            G,
            f"{b}_g_bess",
            "Batarya deşarj geliri",
            "USD",
            lambda j: (
                f"=IF({A['yek']}=1,(1-{phi(j)})*{A['L']}*{self.r(U, f'{b}_gelir_d_q', j)}+{phi(j)}*{pf(j)}*{self.r(U, f'{b}_E_d', j)},{A['L']}*{self.r(U, f'{b}_gelir_d_q', j)})"
            ),
            not_="YEKDEM yılında depolanıp verilen santral enerjisi YEKDEM fiyatından",
        )
        self.seri(
            G,
            f"{b}_g_sfk",
            "SFK kapasite geliri",
            "USD",
            lambda j: f"=IF({self.r(Z, 'pg_ok', j)}=1,INDEX({A['pg_a']},{self.r(Z, 'k', j)})*{self.r(U, f'{b}_yh_a', j)},0)",
        )
        self.seri(G, f"{b}_g_yekg", "YEK-G geliri", "USD", lambda j: f"={A['yekg']}*{self.r(U, f'{b}_satilan', j)}")
        self.seri(
            G,
            f"{b}_azaltim",
            "Batarya dengesizlik azaltımı",
            "oran",
            lambda j: (
                f"=IF({self.r(Z, 'pg_ok', j)}=1,{A['dz_azami']}*MIN(MAX({A['bess_mw']}-{self.r(U, f'{b}_rezerv', j)}/8760,0)/({A['dz_doyum']}*{A['res_mw']}),1),0)"
            ),
            toplam=False,
            fmt="0.000",
        )
        self.seri(
            G, f"{b}_esz", "Eşzamanlılık düzeltmesi", "USD", lambda j: f"={A['esz']}*{A['L']}*{self.r(U, f'{b}_E_res_net', j)}*(1-{phi(j)})"
        )
        self.seri(
            G,
            f"{b}_deng",
            "Dengesizlik maliyeti",
            "USD",
            lambda j: f"={A['deng']}*({self.r(G, f'{b}_g_res', j)}-{self.r(G, f'{b}_esz', j)})*(1-{self.r(G, f'{b}_azaltim', j)})",
        )
        self.seri(
            G,
            f"{b}_g_top",
            "Toplam gelir",
            "USD",
            lambda j: (
                f"={self.r(G, f'{b}_g_res', j)}+{self.r(G, f'{b}_g_bess', j)}+{self.r(G, f'{b}_g_sfk', j)}+{self.r(G, f'{b}_g_yekg', j)}"
            ),
        )
        # OPEX
        isl = lambda j: f"IF({self.r(Z, 'k', j)}>0,1,0)"
        self.seri(
            O, f"{b}_o_sabit", "RES sabit O&M + arazi", "USD", lambda j: f"={isl(j)}*({A['opex_res']}+{A['arazi_res']})*{A['res_mw']}"
        )
        self.seri(O, f"{b}_o_bess", "Batarya O&M", "USD", lambda j: f"={isl(j)}*{A['opex_bess']}*{A['capex_bess_t']}")
        self.seri(O, f"{b}_o_sig", "Sigorta", "USD", lambda j: f"={isl(j)}*{A['sigorta']}*{A['capex']}")
        self.seri(
            O,
            f"{b}_o_teias",
            "TEİAŞ bedelleri",
            "USD",
            lambda j: f"={isl(j)}*{A['teias_k']}*{A['sebeke_mw']}+{A['teias_e']}*{self.r(U, f'{b}_satilan', j)}",
        )
        self.seri(O, f"{b}_o_yip", "Batarya yıpranması", "USD", lambda j: f"={A['yipranma']}*{self.r(U, f'{b}_E_d', j)}")
        self.seri(
            O,
            f"{b}_opex",
            "Toplam OPEX",
            "USD",
            lambda j: "=" + "+".join(self.r(O, f"{b}_{x}", j) for x in ["o_sabit", "o_bess", "o_sig", "o_teias", "o_yip"]),
        )
        self.seri(
            O,
            f"{b}_favok",
            "FAVÖK",
            "USD",
            lambda j: f"={self.r(G, f'{b}_g_top', j)}-{self.r(G, f'{b}_deng', j)}-{self.r(G, f'{b}_esz', j)}-{self.r(O, f'{b}_opex', j)}",
        )
        kk = lambda j: self.r(Z, "k", j)
        self.seri(
            O,
            f"{b}_aug",
            "Augmentation + batarya yenilemesi",
            "USD",
            lambda j: (
                f"=IF(OR({kk(j)}=8,{kk(j)}=16,{kk(j)}=28),{A['aug_oran']}*{A['bess_mwh']}*{A['capex_bess_e']}*(1-{A['aug_dusus']})^{kk(j)},0)"
                f"+IF({kk(j)}={A['omur_bess']},{A['capex_bess_t']}*(1-{A['aug_dusus']})^{A['omur_bess']},0)"
            ),
        )
        self.seri(
            O,
            f"{b}_sokum",
            "Söküm (son işletme yılı)",
            "USD",
            lambda j: f"=IF({kk(j)}={A['omur']},{A['sokum_bess']}*{A['bess_mwh']}+{A['sokum_res']}*{A['res_mw']},0)",
        )
        # Vergi: amortisman, FIFO zarar mahsubu (5 yıl)
        self.seri(
            V,
            f"{b}_amort",
            "Amortisman (10 yıl doğrusal)",
            "USD",
            lambda j: f"=IF(AND({kk(j)}>=1,{kk(j)}<={A['amort']}),{A['capex']}/{A['amort']},0)",
        )
        self.seri(
            V,
            f"{b}_matrah",
            "Mahsup öncesi matrah",
            "USD",
            lambda j: f"={self.r(O, f'{b}_favok', j)}-{self.r(V, f'{b}_amort', j)}-{self.r(O, f'{b}_aug', j)}-{self.r(O, f'{b}_sokum', j)}",
        )
        self.seri(V, f"{b}_zarar", "Dönem zararı", "USD", lambda j: f"=MAX(-{self.r(V, f'{b}_matrah', j)},0)")
        self.seri(V, f"{b}_pozitif", "Pozitif matrah", "USD", lambda j: f"=MAX({self.r(V, f'{b}_matrah', j)},0)")
        w = self.ws(V)
        rw0 = w.max_row + 1
        w.cell(
            rw0, 1, "FIFO mahsup matrisi: satır = zararın doğduğu işletme yılı v, hücre = o yıl kullanılan tutar (v+1 … v+5)"
        ).font = Font(italic=True)
        satirlar = {}
        for vv in range(1, 31):
            rw = rw0 + vv
            satirlar[vv] = rw
            w.cell(rw, 1, f"Zarar yılı {vv} — kullanım")
            w.cell(rw, 2, "USD")
            Lv = f"INDEX({self.aralik(V, f'{b}_zarar')},MATCH({vv},{self.aralik(Z, 'k')},0))"
            for j in range(1, NYIL + 1):
                kc = self.r(Z, "k", j)
                onceki = f"SUM({self.kol(1)}{rw}:{self.kol(j - 1)}{rw})" if j > 1 else "0"
                eski = f"SUM({self.kol(j)}{rw0 + 1}:{self.kol(j)}{rw - 1})" if vv > 1 else "0"
                w.cell(
                    rw, self.C0 + j, f"=IF(AND({kc}>{vv},{kc}-{vv}<=5),MIN({Lv}-{onceki},MAX(0,{self.r(V, f'{b}_pozitif', j)}-{eski})),0)"
                ).number_format = "#,##0"
        self.satir[(V, f"{b}_fifo_bas")] = rw0 + 1
        self.satir[(V, f"{b}_fifo_son")] = rw0 + 30
        self.seri(V, f"{b}_mahsup", "Kullanılan zarar mahsubu", "USD", lambda j: f"=SUM({self.kol(j)}{rw0 + 1}:{self.kol(j)}{rw0 + 30})")
        self.seri(
            V,
            f"{b}_vergi",
            "Kurumlar vergisi",
            "USD",
            lambda j: f"={A['vergi']}*({self.r(V, f'{b}_pozitif', j)}-{self.r(V, f'{b}_mahsup', j)})",
        )
        self.seri(
            V,
            f"{b}_cfads",
            "CFADS (borç servisine esas nakit)",
            "USD",
            lambda j: f"={self.r(O, f'{b}_favok', j)}-{self.r(V, f'{b}_vergi', j)}-{self.r(O, f'{b}_aug', j)}-{self.r(O, f'{b}_sokum', j)}",
        )

    def kdv_borc_nakit(self):
        A, Z, V, B, N_ = self.ad, "Zaman", "Vergi", "Borc", "Nakit_akisi"
        # KDV (teşvik yoksa)
        w = self.ws(V)
        rw = w.max_row + 2
        w.cell(rw, 1, "— KDV (teşvik belgesi yoksa) —").font = Font(bold=True, color="1D4F91")
        tes = f'IF({A["tesvik"]}="Yok",1,0)'
        self.seri(
            V,
            "kdv_odenen",
            "İnşaatta ödenen KDV",
            "USD",
            lambda j: f"={tes}*{A['kdv']}*{A['kdv_kapsam']}*{A['capex']}*{self.r(Z, 'insaat_pay', j)}",
        )
        self.seri(
            V,
            "kdv_gelir",
            "KDV mahsubuna esas gelir",
            "USD",
            lambda j: (
                f"=Gelir!{self.kol(j)}{self.satir[('Gelir', 'Sec_g_res')]}+Gelir!{self.kol(j)}{self.satir[('Gelir', 'Sec_g_bess')]}+Gelir!{self.kol(j)}{self.satir[('Gelir', 'Sec_g_sfk')]}"
            ),
        )
        self.satir[(V, "kdv_bakiye")] = self.ws(V).max_row + 3  # ileri başvuru: on → mahsup → bakiye
        self.seri(
            V,
            "kdv_on",
            "Devreden KDV (aşınma sonrası, mahsup öncesi)",
            "USD",
            lambda j: (
                f"={self.r(V, 'kdv_bakiye', j - 1)}*(1-{A['kdv_asinma']})+{self.r(V, 'kdv_odenen', j)}"
                if j > 1
                else f"={self.r(V, 'kdv_odenen', j)}"
            ),
            toplam=False,
        )
        self.seri(
            V,
            "kdv_mahsup",
            "KDV mahsubu (çıkış KDV'si)",
            "USD",
            lambda j: f"=IF({self.r(Z, 'k', j)}>0,MIN({self.r(V, 'kdv_on', j)},{A['kdv']}*{self.r(V, 'kdv_gelir', j)}),0)",
        )
        self.seri(
            V,
            "kdv_bakiye",
            "Devreden KDV bakiyesi",
            "USD",
            lambda j: f"={self.r(V, 'kdv_on', j)}-{self.r(V, 'kdv_mahsup', j)}",
            toplam=False,
        )
        # Borç
        w = self.ws(B)
        w["A1"] = "Borç: DSCR hedefli heykelsi boyutlandırma, geri ödeme, DSCR ve LLCR"
        w["A1"].font = Font(bold=True, size=13, color="1D4F91")
        self.zaman_basligi(w)
        k = lambda j: self.r(Z, "k", j)
        self.seri(
            B,
            "pay",
            "Sözleşmeli CFADS payı",
            "oran",
            lambda j: (
                f"=IF({k(j)}>0,MIN(MAX({self.r(Z, 'phi', j)}*{self.r(Z, 'pf', j)}*Uretim!{self.kol(j)}{self.satir[('Uretim', 'Sec_E_res_net')]}/MAX(OPEX!{self.kol(j)}{self.satir[('OPEX', 'Sec_favok')]},1),0),1),0)"
            ),
            toplam=False,
            fmt="0.000",
        )
        self.seri(
            B,
            "dscr_h",
            "Hedef DSCR (harmanlanmış)",
            "kat",
            lambda j: f"=IF({k(j)}>0,1/({self.r(B, 'pay', j)}/{A['dscr_s']}+(1-{self.r(B, 'pay', j)})/{A['dscr_t']}),0)",
            toplam=False,
            fmt="0.000",
        )
        self.seri(
            B,
            "servis_ham",
            "Borç servisi (tavan öncesi, nominal)",
            "USD",
            lambda j: (
                f"=IF(AND({k(j)}>=1,{k(j)}<={A['vade']}),MAX(Vergi!{self.kol(j)}{self.satir[('Vergi', 'P90_cfads')]},0)*{self.r(Z, 'enf', j)}/{self.r(B, 'dscr_h', j)},0)"
            ),
        )
        self.seri(
            B,
            "servis_iskonto",
            "Servisin bugünkü değeri (nominal faizle, k'ya göre)",
            "USD",
            lambda j: f"=IF({k(j)}>0,{self.r(B, 'servis_ham', j)}/(1+{A['faiz']})^{k(j)},0)",
        )
        w = self.ws(B)
        rw = w.max_row + 1
        w.cell(rw, 1, "Borç (tavan öncesi, nominal)")
        w.cell(rw, 4, f"=SUM({self.aralik(B, 'servis_iskonto')})").number_format = "#,##0"
        self.ad["borc_ham"] = f"Borc!$D${rw}"
        rw += 1
        w.cell(rw, 1, "Ölçek (azami borç/CAPEX tavanı)")
        w.cell(rw, 4, f"=MIN(1,{A['azami_borc']}*{A['capex_nom']}/MAX({A['borc_ham']},1))").number_format = "0.0000"
        self.ad["olcek"] = f"Borc!$D${rw}"
        rw += 1
        w.cell(rw, 1, "Borç tutarı (nominal)")
        w.cell(rw, 4, f"={A['borc_ham']}*{A['olcek']}").number_format = "#,##0"
        self.ad["borc_nom"] = f"Borc!$D${rw}"
        rw += 1
        w.cell(rw, 1, "Borç tutarı (reel 2026)")
        w.cell(rw, 4, f"={A['borc_nom']}/(1+{A['pi']})^2").number_format = "#,##0"
        self.ad["borc_reel"] = f"Borc!$D${rw}"
        rw += 1
        self.seri(B, "servis", "Borç servisi (nominal)", "USD", lambda j: f"={self.r(B, 'servis_ham', j)}*{A['olcek']}")
        self.satir[(B, "kapanis")] = self.ws(B).max_row + 3  # ileri başvuru: açılış → faiz → kapanış
        self.seri(
            B,
            "acilis",
            "Açılış bakiyesi",
            "USD",
            lambda j: (
                f"=IF({k(j)}=1,{A['borc_nom']},IF({k(j)}>1,{self.r(B, 'kapanis', j - 1)},0))"
                if j > 1
                else f"=IF({k(j)}=1,{A['borc_nom']},0)"
            ),
            toplam=False,
        )
        self.seri(B, "faiz", "Faiz", "USD", lambda j: f"={self.r(B, 'acilis', j)}*{A['faiz']}")
        self.seri(
            B,
            "kapanis",
            "Kapanış bakiyesi",
            "USD",
            lambda j: f"={self.r(B, 'acilis', j)}+{self.r(B, 'faiz', j)}-{self.r(B, 'servis', j)}",
            toplam=False,
        )
        cf_sec = lambda j: f"Vergi!{self.kol(j)}{self.satir[('Vergi', 'Sec_cfads')]}"
        self.seri(
            B,
            "dscr",
            "DSCR",
            "kat",
            lambda j: f'=IF({self.r(B, "servis", j)}>0,{cf_sec(j)}*{self.r(Z, "enf", j)}/{self.r(B, "servis", j)},"")',
            toplam=False,
            fmt="0.00",
        )
        self.seri(
            B,
            "cfads_nom",
            "CFADS (nominal, vade içinde)",
            "USD",
            lambda j: f"=IF(AND({k(j)}>=1,{k(j)}<={A['vade']}),{cf_sec(j)}*{self.r(Z, 'enf', j)},0)",
        )
        self.seri(
            B,
            "pv_ileri",
            "Kalan CFADS'nin bugünkü değeri (faizle)",
            "USD",
            lambda j: (
                f"=IF({self.r(B, 'cfads_nom', j)}<>0,({self.r(B, 'cfads_nom', j)}+{self.r(B, 'pv_ileri', j + 1)})/(1+{A['faiz']}),0)"
                if j < NYIL
                else f"=IF({self.r(B, 'cfads_nom', j)}<>0,{self.r(B, 'cfads_nom', j)}/(1+{A['faiz']}),0)"
            ),
            toplam=False,
        )
        self.seri(
            B,
            "llcr",
            "LLCR",
            "kat",
            lambda j: f'=IF({self.r(B, "acilis", j)}>1,{self.r(B, "pv_ileri", j)}/{self.r(B, "acilis", j)},"")',
            toplam=False,
            fmt="0.00",
        )
        self.seri(
            B,
            "kalkan",
            "Faiz vergi kalkanı (reel)",
            "USD",
            lambda j: (
                f"=IF({k(j)}>0,MIN({self.r(B, 'faiz', j)}*{A['vergi']},MAX(Vergi!{self.kol(j)}{self.satir[('Vergi', 'Sec_vergi')]},0)*{self.r(Z, 'enf', j)})/{self.r(Z, 'enf', j)},0)"
            ),
        )
        # Nakit akışı
        w = self.ws(N_)
        w["A1"] = "Proje ve öz kaynak nakit akışları (reel 2026 USD)"
        w["A1"].font = Font(bold=True, size=13, color="1D4F91")
        self.zaman_basligi(w)
        prim = f"-{A['prim']}"
        self.seri(N_, "capex", "CAPEX", "USD", lambda j: f"=-{A['capex']}*{self.r(Z, 'insaat_pay', j)}")
        self.seri(N_, "kdv", "KDV (ödenen − mahsup)", "USD", lambda j: f"=-{self.r(V, 'kdv_odenen', j)}+{self.r(V, 'kdv_mahsup', j)}")
        self.seri(N_, "cfads", "CFADS", "USD", lambda j: f"={cf_sec(j)}")
        self.seri(
            N_,
            "proje",
            "Proje nakit akışı",
            "USD",
            lambda j: f"={self.r(N_, 'capex', j)}+{self.r(N_, 'kdv', j)}+{self.r(N_, 'cfads', j)}",
            deger_t0=f"={prim}",
        )
        self.seri(N_, "borc_cek", "Borç kullanımı (reel)", "USD", lambda j: f"={A['borc_reel']}*{self.r(Z, 'insaat_pay', j)}")
        self.seri(
            N_,
            "borc_odeme",
            "Borç servisi net (reel, kalkan sonrası)",
            "USD",
            lambda j: f"=IF({self.r(Z, 'k', j)}>0,-({self.r(B, 'servis', j)}/{self.r(Z, 'enf', j)}-{self.r(B, 'kalkan', j)}),0)",
        )
        self.seri(
            N_,
            "oz",
            "Öz kaynak nakit akışı",
            "USD",
            lambda j: f"={self.r(N_, 'proje', j)}+{self.r(N_, 'borc_cek', j)}+{self.r(N_, 'borc_odeme', j)}",
            deger_t0=f"={prim}",
        )
        self.seri(
            N_,
            "kum",
            "Kümülatif proje nakit akışı",
            "USD",
            lambda j: (
                f"={self.r(N_, 'kum', j - 1)}+{self.r(N_, 'proje', j)}" if j > 1 else f"={self.r(N_, 'kum', 0)}+{self.r(N_, 'proje', j)}"
            ),
            toplam=False,
            deger_t0=f"={self.r(N_, 'proje', 0)}",
        )
        self.seri(
            N_,
            "geri_odendi",
            "Geri ödendi mi (işletmede kümülatif ≥ 0)",
            "0/1",
            lambda j: f"=IF(AND({self.r(Z, 'k', j)}>0,{self.r(N_, 'kum', j)}>=0),1,0)",
            fmt="0",
            toplam=False,
        )
        G_ = "Gelir"
        self.seri(
            N_,
            "maliyet",
            "LCOE maliyeti (gelir − vergi − CFADS)",
            "USD",
            lambda j: f"={self.r(G_, 'Sec_g_top', j)}-Vergi!{self.kol(j)}{self.satir[('Vergi', 'Sec_vergi')]}-{cf_sec(j)}",
        )

    def ozet(self, mc):
        A, N_, B, Z = self.ad, "Nakit_akisi", "Borc", "Zaman"
        w = self.ws("Ozet")
        w["A1"] = "Özet — seçilen senaryo"
        w["A1"].font = Font(bold=True, size=14, color="1D4F91")
        w.column_dimensions["A"].width = 52
        w.column_dimensions["B"].width = 20
        w.column_dimensions["C"].width = 70
        satirlar = [
            ("Senaryo", f"={A['anahtar']}", "@", ""),
            (
                "NPV (edinim primi sonrası) [M USD]",
                f"=SUMPRODUCT({self.aralik(N_, 'proje', 0, NYIL)},{self.aralik(Z, 'df', 0, NYIL)})/1e6",
                "0.00",
                "%7 reel WACC, t = 0 bugün",
            ),
            (
                "Proje IRR (reel)",
                f"=IFERROR(IRR({self.aralik(N_, 'proje', 0, NYIL)},0.05),IRR({self.aralik(N_, 'proje', 0, NYIL)},-0.05))",
                "0.00%",
                "",
            ),
            (
                "Öz kaynak IRR (reel)",
                f'=IFERROR(IRR({self.aralik(N_, "oz", 0, NYIL)},0.05),IFERROR(IRR({self.aralik(N_, "oz", 0, NYIL)},-0.05),"tanımsız"))',
                "0.00%",
                "Kök yoksa 'tanımsız'",
            ),
            (
                "LCOE (satılan enerji) [USD/MWh]",
                f"=(SUMPRODUCT(-{self.aralik(N_, 'capex')},{self.aralik(Z, 'dfl')})+SUMPRODUCT({self.aralik(N_, 'maliyet')},{self.aralik(Z, 'dfl')}))/SUMPRODUCT(Uretim!${self.kol(1)}${self.satir[('Uretim', 'Sec_satilan')]}:${self.kol(NYIL)}${self.satir[('Uretim', 'Sec_satilan')]},{self.aralik(Z, 'dfl')})",
                "0.00",
                "Vergi öncesi, inşaat başından iskonto",
            ),
            ("Asgari DSCR", f"=MIN({self.aralik(B, 'dscr')})", "0.00", "Borç vadesi boyunca"),
            ("Ortalama DSCR", f"=AVERAGE({self.aralik(B, 'dscr')})", "0.00", ""),
            ("Asgari LLCR", f"=MIN({self.aralik(B, 'llcr')})", "0.00", ""),
            ("Borç (reel) [M USD]", f"={A['borc_reel']}/1e6", "0.00", ""),
            (
                "Geri ödeme süresi (COD'dan, yıl)",
                f'=IFERROR(INDEX({self.aralik(Z, "k")},MATCH(1,{self.aralik(N_, "geri_odendi")},0)),"> ömür")',
                "0",
                "İskontosuz proje nakit akışı",
            ),
            ("Azami edinim primi [M USD]", f"=MAX(B4+{A['prim']}/1e6,0)", "0.00", "Primsiz NPV (pozitifse)"),
            ("YEKDEM'e katılınan yıl sayısı", f"=SUM({self.aralik(Z, 'katil')})", "0", "Deterministik fiyatta karar kuralı"),
        ]
        for i, (e, f, fmt, n) in enumerate(satirlar):
            rw = 3 + i
            w.cell(rw, 1, e)
            c = w.cell(rw, 2, f)
            c.number_format = fmt
            w.cell(rw, 3, n)
            self.ad["oz_" + str(i)] = f"Ozet!$B${rw}"
        rw = 3 + len(satirlar) + 2
        w.cell(rw, 1, "Python Monte Carlo referansı (deterministik model değil; outputs/07/baz_vaka.csv)").font = Font(
            bold=True, color="1D4F91"
        )
        for _, rr in mc.iterrows():
            rw += 1
            w.cell(rw, 1, rr["Senaryo"])
            w.cell(rw, 2, round(float(rr["Beklenen NPV [M USD]"]), 2))
            w.cell(
                rw,
                3,
                f"P(NPV>0) %{rr['P(NPV > 0) [%]']:.0f} · CVaR %5 {rr['CVaR %5 [M USD]']:.1f} M$ · azami prim {rr['Azami edinim primi [M USD]']:.2f} M$",
            )
        ay = CIKTILAR / "07" / "fark_ayristirma_sonra.csv"
        if ay.exists():  # deterministik baz → Monte Carlo beklenen değer köprüsü (Bölüm A)
            rw += 2
            w.cell(rw, 1, "Deterministik → Monte Carlo köprüsü (outputs/07/fark_ayristirma_sonra.csv)").font = Font(
                bold=True, color="1D4F91"
            )
            for _, rr in pd.read_csv(ay).iterrows():
                rw += 1
                w.cell(rw, 1, rr["Adım"])
                w.cell(rw, 2, round(float(rr["E[NPV] katkısı [M USD]"]), 2))
                w.cell(rw, 3, rr["Bileşen"])

    def senaryolar_ve_kontroller(self, py: pd.DataFrame, ex: pd.DataFrame | None):
        A = self.ad
        w = self.ws("Senaryolar")
        w["A1"] = "Senaryolar — Python ve Excel (LibreOffice ile yeniden hesaplanmış) sonuçları"
        w["A1"].font = Font(bold=True, size=13, color="1D4F91")
        w["A2"] = (
            "Python: src/risk/degerleme.nakit_akislari aynı deterministik girdilerle. Excel: bu dosya her senaryo için LibreOffice'te yeniden hesaplandı."
        )
        kol = [
            "Anahtar",
            "Gelir yapısı",
            "Fiyat",
            "Üretim",
            "Teşvik",
            "Python NPV [M USD]",
            "Python proje IRR [%]",
            "Python öz kaynak IRR [%]",
            "Python asgari DSCR",
            "Python LCOE",
            "Excel NPV [M USD]",
            "Excel proje IRR [%]",
            "Excel öz kaynak IRR [%]",
            "Excel asgari DSCR",
            "Excel LCOE",
            "Fark NPV [M USD]",
            "Fark proje IRR [puan]",
            "Fark öz kaynak IRR [puan]",
            "Durum",
        ]
        for i, h in enumerate(kol):
            c = w.cell(4, i + 1, h)
            c.font = BASLIK
            c.fill = BASLIK_DOLGU
            c.alignment = Alignment(wrap_text=True)
            w.column_dimensions[CL(i + 1)].width = 14
        w.column_dimensions["A"].width = 34
        for r_i, (_, rr) in enumerate(py.iterrows()):
            rw = 5 + r_i
            vals = [
                rr["Anahtar"],
                rr["yapi"],
                rr["seviye"],
                rr["durum"],
                rr["tesvik"],
                rr["NPV [M USD]"],
                rr["Proje IRR [%]"],
                rr["Öz kaynak IRR [%]"],
                rr["Asgari DSCR"],
                rr["LCOE [USD/MWh]"],
            ]
            for c_i, v_ in enumerate(vals):
                if isinstance(v_, float) and v_ != v_:
                    v_ = "tanımsız"
                x = w.cell(rw, c_i + 1, v_)
                x.number_format = "0.000" if isinstance(v_, float) else "@"
            if ex is not None:
                e = ex.loc[rr["Anahtar"]]
                for c_i, k_ in enumerate(["NPV [M USD]", "Proje IRR [%]", "Öz kaynak IRR [%]", "Asgari DSCR", "LCOE [USD/MWh]"]):
                    x = float(e[k_])
                    w.cell(rw, 11 + c_i, x if x == x else "tanımsız").number_format = "0.000"
                w.cell(rw, 16, f"=K{rw}-F{rw}").number_format = "0.0000"
                w.cell(rw, 17, f"=L{rw}-G{rw}").number_format = "0.0000"
                w.cell(rw, 18, f"=M{rw}-H{rw}").number_format = "0.0000"
                w.cell(rw, 19, f'=IF(AND(ABS(P{rw})<0.01,IFERROR(ABS(Q{rw})<0.01,TRUE),IFERROR(ABS(R{rw})<0.01,TRUE)),"Mutabık","FARK")')
        self.ad["sen_ilk"], self.ad["sen_son"] = 5, 4 + len(py)
        # Kontroller
        w = self.ws("Kontroller")
        w["A1"] = "Kontroller"
        w["A1"].font = Font(bold=True, size=14, color="1D4F91")
        w.column_dimensions["A"].width = 70
        w.column_dimensions["B"].width = 18
        w.column_dimensions["C"].width = 14
        N_, B, Z = "Nakit_akisi", "Borc", "Zaman"
        s0, s1 = self.ad["sen_ilk"], self.ad["sen_son"]
        bul = lambda kc: f"INDEX(Senaryolar!${kc}${s0}:${kc}${s1},MATCH({A['anahtar']},Senaryolar!$A${s0}:$A${s1},0))"
        hata_say = "+".join(
            f"SUMPRODUCT(--ISERROR('{s}'!A1:{CL(self.son)}{self.wb[s].max_row}))"
            for s in ["Zaman", "Uretim", "Gelir", "OPEX", "Vergi", "Borc", "Nakit_akisi"]
        )
        kontroller = [
            (
                "Nakit tutarlılığı: Σ proje akışı = Σ CAPEX + Σ KDV + Σ CFADS − prim [USD]",
                f"=SUM({self.aralik(N_, 'proje', 0, NYIL)})-(SUM({self.aralik(N_, 'capex')})+SUM({self.aralik(N_, 'kdv')})+SUM({self.aralik(N_, 'cfads')})-{A['prim']})",
                "abs<1",
            ),
            (
                "CFADS = FAVÖK − vergi − augmentation − söküm (Σ fark) [USD]",
                f"=SUM(Vergi!{self.kol(1)}{self.satir[('Vergi', 'Sec_cfads')]}:{self.kol(NYIL)}{self.satir[('Vergi', 'Sec_cfads')]})-(SUM(OPEX!{self.kol(1)}{self.satir[('OPEX', 'Sec_favok')]}:{self.kol(NYIL)}{self.satir[('OPEX', 'Sec_favok')]})-SUM(Vergi!{self.kol(1)}{self.satir[('Vergi', 'Sec_vergi')]}:{self.kol(NYIL)}{self.satir[('Vergi', 'Sec_vergi')]})-SUM(OPEX!{self.kol(1)}{self.satir[('OPEX', 'Sec_aug')]}:{self.kol(NYIL)}{self.satir[('OPEX', 'Sec_aug')]})-SUM(OPEX!{self.kol(1)}{self.satir[('OPEX', 'Sec_sokum')]}:{self.kol(NYIL)}{self.satir[('OPEX', 'Sec_sokum')]}))",
                "abs<1",
            ),
            (
                "Borç bakiyesi vade sonunda sıfır (vade yılındaki kapanış) [USD]",
                f"=INDEX({self.aralik(B, 'kapanis')},MATCH({A['vade']},{self.aralik(Z, 'k')},0))",
                "abs<1",
            ),
            (
                "Negatif vergi olan yıl sayısı",
                f'=COUNTIF(Vergi!{self.kol(1)}{self.satir[("Vergi", "Sec_vergi")]}:{self.kol(NYIL)}{self.satir[("Vergi", "Sec_vergi")]},"<-0.001")',
                "=0",
            ),
            ("Python mutabakatı — NPV farkı [M USD]", f"={A['oz_1']}-{bul('F')}", "abs<0.01"),
            ("Python mutabakatı — proje IRR farkı [puan]", f"=100*{A['oz_2']}-{bul('G')}", "abs<0.01"),
            ("Python mutabakatı — öz kaynak IRR farkı [puan]", f"=100*{A['oz_3']}-{bul('H')}", "abs<0.01"),
            ("Python mutabakatı — asgari DSCR farkı", f"={A['oz_5']}-{bul('I')}", "abs<0.001"),
            ("Python mutabakatı — LCOE farkı [USD/MWh]", f"={A['oz_4']}-{bul('J')}", "abs<0.01"),
            ("Formül hatası sayısı (#REF!, #DIV/0!, #VALUE!, #NAME?, #N/A)", f"={hata_say}", "=0"),
            (
                "Tüm senaryolarda mutabakat (Senaryolar sayfası)",
                f'=COUNTIF(Senaryolar!$S${s0}:$S${s1},"Mutabık")&" / "&{s1 - s0 + 1}',
                "bilgi",
            ),
        ]
        for i, h in enumerate(["Kontrol", "Değer", "Ölçüt", "Durum"]):
            c = w.cell(3, i + 1, h)
            c.font = BASLIK
            c.fill = BASLIK_DOLGU
        for i, (e, f, olc) in enumerate(kontroller):
            rw = 4 + i
            w.cell(rw, 1, e)
            w.cell(rw, 2, f).number_format = "0.0000"
            w.cell(rw, 3, olc)
            if olc.startswith("abs<"):
                w.cell(rw, 4, f'=IF(ABS(B{rw})<{olc[4:]},"OK","HATA")')
            elif olc == "=0":
                w.cell(rw, 4, f'=IF(B{rw}=0,"OK","HATA")')

    def kur(self, py, ex, mc):
        from openpyxl.workbook.properties import CalcProperties

        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        self.kapak()
        self.girdiler()
        self.python_girdileri()
        self.zaman()
        self.blok("Sec")
        self.blok("P90")
        self.kdv_borc_nakit()
        self.ozet(mc)
        self.senaryolar_ve_kontroller(py, ex)
        for s in self.wb.sheetnames:
            for row in self.wb[s].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and c.font != MAVI:
                        pass
        return self.wb


# =====================================================================================================================
# 3. LibreOffice ile yeniden hesaplama ve mutabakat
# =====================================================================================================================
def lo_hesapla(yol: Path, cikti_dizin: Path) -> Path:
    subprocess.run(
        [str(SOFFICE), "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(cikti_dizin), str(yol)],
        check=True,
        capture_output=True,
    )
    return cikti_dizin / yol.name


def excel_oku(yol: Path) -> dict:
    wb = openpyxl.load_workbook(yol, data_only=True)
    o = wb["Ozet"]

    def f(r):
        x = o.cell(r, 2).value
        return float(x) if isinstance(x, (int, float)) else float("nan")

    return {
        "NPV [M USD]": f(4),
        "Proje IRR [%]": 100 * f(5),
        "Öz kaynak IRR [%]": 100 * f(6),
        "Asgari DSCR": f(8),
        "LCOE [USD/MWh]": f(7),
        "hata": wb["Kontroller"]["B13"].value,
    }


def main():
    CIKTI.mkdir(parents=True, exist_ok=True)
    pm = PythonModel()
    mc = pd.read_csv(CIKTILAR / "07" / "baz_vaka.csv")
    satir = []
    for yapi, seviye, durum, tes in itertools.product(YAPILAR, SEVIYE, ["P50", "P90"], ["Var", "Yok"]):
        r = pm.sonuc(yapi, seviye, durum, tes == "Var")
        satir.append({"Anahtar": f"{yapi}|{seviye}|{durum}|{tes}", "yapi": yapi, "seviye": seviye, "durum": durum, "tesvik": tes, **r})
    py = pd.DataFrame(satir)
    print(
        py[["Anahtar", "NPV [M USD]", "Proje IRR [%]", "Öz kaynak IRR [%]", "Asgari DSCR", "LCOE [USD/MWh]", "YEKDEM katılım yılı"]]
        .round(3)
        .head(16)
        .to_string(index=False)
    )
    # 1. geçiş: Excel'i kur, her senaryo için LibreOffice'te hesaplat
    wb = Kitap(pm).kur(py, None, mc)
    wb.save(DOSYA)
    tmp = Path(tempfile.mkdtemp())
    ex = {}
    for _, rr in py.iterrows():
        wbs = openpyxl.load_workbook(DOSYA)
        g = wbs["Girdiler"]
        g["C6"], g["C7"], g["C8"], g["C9"] = rr["yapi"], rr["seviye"], rr["durum"], rr["tesvik"]
        gec = tmp / "giris"
        gec.mkdir(exist_ok=True)
        cik = tmp / "cikis"
        cik.mkdir(exist_ok=True)
        yol = gec / "m.xlsx"
        wbs.save(yol)
        ex[rr["Anahtar"]] = excel_oku(lo_hesapla(yol, cik))
    ex = pd.DataFrame(ex).T
    # 2. geçiş: sonuçları Senaryolar sayfasına yaz, baz senaryoyla kaydet ve LibreOffice ile son kez hesapla
    wb = Kitap(pm).kur(py, ex, mc)
    wb.save(DOSYA)
    son = lo_hesapla(DOSYA, tmp)
    wbk = openpyxl.load_workbook(son, data_only=True)
    kont = [
        (wbk["Kontroller"].cell(r, 1).value, wbk["Kontroller"].cell(r, 2).value, wbk["Kontroller"].cell(r, 4).value) for r in range(4, 15)
    ]
    mut = py.set_index("Anahtar")[["NPV [M USD]", "Proje IRR [%]", "Öz kaynak IRR [%]", "Asgari DSCR", "LCOE [USD/MWh]"]].join(
        ex[["NPV [M USD]", "Proje IRR [%]", "Öz kaynak IRR [%]", "Asgari DSCR", "LCOE [USD/MWh]", "hata"]], rsuffix=" (Excel)"
    )
    for k_ in ["NPV [M USD]", "Proje IRR [%]", "Öz kaynak IRR [%]", "Asgari DSCR", "LCOE [USD/MWh]"]:
        mut[f"Fark {k_}"] = mut[f"{k_} (Excel)"].astype(float) - mut[k_]
    mut.to_csv(CIKTI / "excel_mutabakat.csv", encoding="utf-8-sig", float_format="%.5f")
    print("\nMutabakat (azami mutlak fark):")
    print(mut[[c for c in mut.columns if c.startswith("Fark")]].abs().max().round(6).to_string())
    print("Formül hatası (her senaryo):", sorted(set(mut["hata"].astype(float))))
    print("\nKontroller (baz senaryo):")
    for k_ in kont:
        print(f"  {k_[2] or '':5} {k_[0]}: {k_[1]}")
    # Hata dizesi taraması
    hatalar = 0
    for s in wbk.sheetnames:
        for row in wbk[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper():
                    hatalar += 1
    print("Yeniden hesaplanmış dosyada hata hücresi:", hatalar)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
