"""
Aşama 7 — 5. Risk kaydı (risk register): 20 risk, olasılık × etki (1-5), modeldeki sayısal etkisi, azaltım ve kalıntı risk.

Puanlar uzman yargısıdır (TAHMİN); sayısal etkiler model çıktılarından alınmıştır (kaynağı "Modelde" sütununda).
Çıktılar: outputs/07/risk_kaydi.xlsx, risk_kaydi.csv, risk_isi_haritasi.png; results/risk_kaydi.csv
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from src import gorsel as g
from src.ayarlar import CIKTILAR, KOK

CIKTI = CIKTILAR / "07"

# (kategori, risk, açıklama, olasılık, etki, sayısal etki, modelde (bölüm), azaltım, kalıntı olasılık, kalıntı etki)
RISKLER = [
    ("Piyasa", "Elektrik fiyat seviyesi düşük kalır", "Uzun dönem PTF beklentinin (58 $/MWh) altında kalır", 4, 5,
     "Varyansın %45–85'i; 45 $/MWh'de NPV −28 M$ (tornado)", "Evet — Aşama 3 tornado, Aşama 5–7 Monte Carlo",
     "PPA veya YEKDEM ile fiyat sabitleme; muhafazakâr borç", 3, 3),
    ("Piyasa", "SFK piyasasının doyumu", "28–33 GW depolama boru hattı ~1 GW SFK ihtiyacını doyurur, fiyat fırsat maliyetine iner", 4, 4,
     "40 MW bataryalı tasarımın E[NPV]'si 28,4 → 6,2 M$; doyum varyansın ~%34'ü", "Evet — Aşama 6 stres testi",
     "30 MW ile başla, büyütmeyi SFK fiyatına bağla; batarya kararını FID'de gözlenen doyumla ver", 4, 2),
    ("Düzenleyici", "TEİAŞ hibrit tesisteki bataryayı SFK'ya kabul etmez", "Tek bağlantı içinde hibrit tesis için emsal yok", 2, 3,
     "Bataryanın değeri yalnız RES'e göre −0,8 M$ (30 MW)", "Evet — Aşama 6 'katılamama' senaryosu",
     "Batarya siparişinden önce yazılı ön yeterlilik", 1, 2),
    ("Düzenleyici / sistem", "Sistem kaynaklı kesinti (YAT talimatları) artar", "Güneş büyümesi, düşük talep ve yüksek hidrolojide RES'lere YAT; 2026'da Ezine RES'lerinde ~%9,5", 4, 3,
     "E[NPV] −2 ile −3 M$; YEKDEM'de kesilen enerjiye YEK farkı ödenmiyor", "Evet — Aşama 7 §1, Ö1",
     "DGP'ye düşük fiyatlı YAT teklifi (tazmin), PPA'da 'deemed generation' maddesi, batarya (%16 kurtarma)", 3, 2),
    ("Düzenleyici", "YEKDEM taban/tavanı nominal kalır", "7189 sayılı karardaki USD sınırları endekslenmezse reel değer 10 yılda ~%22 aşınır", 3, 4,
     "Nihai tasarım, COD 2030: E[NPV] −3,2 (endeksli) vs −11,4 M$ (nominal)", "Evet — Aşama 7 §2, Ö1",
     "FID öncesi hukuki görüş; nominal ise PPA'yı seç", 3, 2),
    ("Düzenleyici", "Azami fiyat tavanı / piyasa müdahalesi", "2022–23 AUF benzeri tavan yenilenebilir gelirini kırpar", 3, 3,
     "AUF dönemi gelir −%44; Monte Carlo'da kriz/tavan olayı (varyansın ~%0,5'i)", "Evet — Aşama 2, Aşama 5 kriz olayı",
     "PPA/YEKDEM (tavandan etkilenmez)", 2, 2),
    ("İnşaat / takvim", "ÇED, imar ve lisans gecikmesi", "Geliştirme aşamasından başlanırsa kritik yol ÇED → imar → lisans", 4, 4,
     "COD P50 Oca 2032 (+48 ay); inşaata hazırda bile 2 yıl kayma tüccar E[NPV]'yi 4,3 → 0,8 M$ indiriyor", "Evet — Aşama 7 §4",
     "İnşaata hazır proje edinimi (COD P50 Mar 2030)", 3, 3),
    ("Düzenleyici / takvim", "YEKDEM için 31.12.2030 COD sınırını kaçırma", "COD 2030 sonrasına kayarsa YEKDEM'e erişim yok", 3, 4,
     "Geliştirme senaryosunda P(COD ≤ 2030) ≈ %0; inşaata hazırda ~%100", "Evet — Aşama 7 §4",
     "İnşaata hazır proje, türbin slot rezervasyonu, gecikme cezalı EPC", 2, 3),
    ("Teknik", "Rüzgâr kaynağı / model hatası", "ERA5 üç santralle kalibre; saha ölçümü ve iz kaybı analizi yok", 3, 3,
     "RES model çarpanı varyansın %3–5'i; 1 yıllık P90 = P50'nin %90'ı", "Evet — Aşama 5 Monte Carlo",
     "12 ay met direği/LiDAR, MCP, WAsP/WindPRO iz analizi", 2, 2),
    ("Tedarik zinciri", "Türbin tedarik gecikmesi ve fiyat artışı", "OEM kapasite darboğazı, Avrupa OEM'lerinin mali zorlukları, navlun", 3, 4,
     "Türbin tedariki kritiklik %98; CAPEX +%20 → NPV −21 M$", "Evet — takvim MC, tornado",
     "Erken slot rezervasyonu, sabit fiyatlı TSA, gecikme cezası (LD), alternatif OEM", 2, 3),
    ("İnşaat", "Şebeke bağlantısı (TM/ENH) gecikmesi", "TEİAŞ iletim yatırımları ve kamulaştırma gecikmesi", 3, 3,
     "Kritiklik indeksi %60; 9–20 ay", "Evet — Aşama 7 §4",
     "Bağlantı anlaşmasını erken imzala; TM/ENH'yi proje şirketi yapsın (TEİAŞ'a devir)", 2, 2),
    ("Teknik", "Batarya bozunması, performans ve yangın güvenliği", "LFP yaşlanması sık döngüde hızlanır; termal kaçak riski", 2, 3,
     "Bozunma varyansın %0,7'si", "Kısmen — Aşama 5 bozunma çarpanı",
     "Performans garantisi, augmentation sözleşmesi, NFPA 855 tasarımı, sigorta", 1, 2),
    ("Çevresel / sosyal", "Ornitoloji: kuş göç yolu", "Çanakkale Boğazı süzülen kuş göç yolu üzerinde: ÇED'de ret/tadil, göç döneminde zorunlu durdurma", 3, 4,
     "Modelde yok — göç dönemi seçici durdurmada tahmini %1–3 üretim kaybı", "Hayır",
     "İki göç dönemi temel çalışma, mikro-yerleşim, radar destekli seçici durdurma, işletme izlemesi", 2, 3),
    ("Çevresel / sosyal", "Yerel kabul ve dava riski", "Tarım/zeytinlik arazisi, görsel etki, gürültü; ÇED kararına dava", 3, 3,
     "Modelde yok (takvimde ÇED/imar kötümser süreleri)", "Kısmen — takvim belirsizliği",
     "Paydaş katılım planı, yerel istihdam, topluluk fayda payı, şikâyet mekanizması", 2, 2),
    ("Finansman", "Kur riski", "PTF ve YEKDEM TL ödenir; TL'nin reel değer kaybı ile USD endeksleme arasındaki fark", 3, 3,
     "Modelde ayrı modellenmedi (reel USD)", "Hayır",
     "USD borç–USD endeksli gelir eşleşmesi, PPA'yı USD'ye endeksle, kur riskinden korunma", 2, 2),
    ("Finansman", "Faiz / sermaye maliyeti artışı", "Küresel faizler veya ülke risk primi yükselir", 3, 4,
     "WACC %8,5'te NPV −15 M$ (tornado)", "Evet — Aşama 3 tornado",
     "DFI finansmanı (EBRD/IFC), faiz swapı, uzun vade", 2, 3),
    ("Finansman", "Teşvik belgesi alınamaz", "KDV istisnası ve gümrük muafiyeti olmadan KDV ön finansmanı", 2, 3,
     "NPV −5,8 M$ (KDV, TL aşınması dahil) + batarya gümrüğü ≤ 0,6 M$", "Evet — Aşama 7 §2",
     "Teşvik belgesini FID öncesi al (9903 sayılı karar)", 1, 2),
    ("Karşı taraf", "PPA alıcısının temerrüdü", "Kurumsal alıcı 12 yıl içinde sözleşmeyi bozar", 2, 3,
     "Modelde yok; temerrütte gelir tüccar yapıya döner", "Hayır",
     "Yatırım notlu alıcı, teminat mektubu, ikame alıcı maddesi", 1, 2),
    ("Piyasa", "Dengesizlik maliyetinin artması", "Yenilenebilir payı arttıkça SMF–PTF farkı büyür", 3, 2,
     "RES %8 dengesizlikte NPV −5 M$ (tornado)", "Evet — tornado, dengesizlik azaltımı",
     "Toplayıcı/portföy, batarya, iyi tahmin", 2, 2),
    ("Karşı taraf", "EPC / OEM iflası veya garanti ihlali", "Yüklenicinin mali zorluğu inşaatı durdurur, garanti boşa düşer", 2, 4,
     "Modelde yok", "Hayır", "Performans ve avans teminatları, ana şirket garantisi, finansal durum taraması", 1, 3),
]


def tablo() -> pd.DataFrame:
    df = pd.DataFrame(RISKLER, columns=["Kategori", "Risk", "Açıklama", "Olasılık (1-5)", "Etki (1-5)", "Sayısal etki",
                                        "Modelde (bölüm)", "Azaltım önlemi", "Kalıntı olasılık", "Kalıntı etki"])
    df.insert(0, "No", [f"R{i + 1:02d}" for i in range(len(df))])
    df.insert(6, "Puan", df["Olasılık (1-5)"] * df["Etki (1-5)"])
    df["Kalıntı puan"] = df["Kalıntı olasılık"] * df["Kalıntı etki"]
    return df


def seviye(p):
    return "Kritik" if p >= 15 else "Yüksek" if p >= 10 else "Orta" if p >= 5 else "Düşük"


def excel(df: pd.DataFrame, yol: Path):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    renk = {"Kritik": "E74C3C", "Yüksek": "F39C12", "Orta": "F7DC6F", "Düşük": "82E0AA"}
    d = df.copy(); d["Seviye"] = d["Puan"].map(seviye); d["Kalıntı seviye"] = d["Kalıntı puan"].map(seviye)
    with pd.ExcelWriter(yol, engine="openpyxl") as w:
        d.to_excel(w, sheet_name="Risk kaydı", index=False)
        ws = w.sheets["Risk kaydı"]
        genislik = {"No": 6, "Kategori": 16, "Risk": 34, "Açıklama": 48, "Sayısal etki": 40, "Modelde (bölüm)": 30, "Azaltım önlemi": 48}
        for j, kol in enumerate(d.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = genislik.get(kol, 11)
            ws.cell(1, j).font = Font(bold=True, color="FFFFFF"); ws.cell(1, j).fill = PatternFill("solid", fgColor="1D4F91")
            ws.cell(1, j).alignment = Alignment(wrap_text=True, vertical="center")
        for i in range(2, len(d) + 2):
            for j in range(1, len(d.columns) + 1):
                ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
            for kol in ("Seviye", "Kalıntı seviye"):
                c = ws.cell(i, list(d.columns).index(kol) + 1)
                c.fill = PatternFill("solid", fgColor=renk[c.value])
        ws.freeze_panes = "C2"


def isi_haritasi(df: pd.DataFrame, yol: Path):
    fig, axs = plt.subplots(1, 2, figsize=(13, 6))
    cm = np.array([[o * e for o in range(1, 6)] for e in range(1, 6)])
    renkler = np.vectorize(lambda p: {"Kritik": "#e74c3c", "Yüksek": "#f39c12", "Orta": "#f7dc6f", "Düşük": "#a9dfbf"}[seviye(p)])(cm)
    for ax, (o_k, e_k, baslik_) in zip(axs, [("Olasılık (1-5)", "Etki (1-5)", "Doğal risk (azaltım öncesi)"),
                                              ("Kalıntı olasılık", "Kalıntı etki", "Kalıntı risk (azaltım sonrası)")]):
        for e in range(5):
            for o in range(5):
                ax.add_patch(plt.Rectangle((o + 0.5, e + 0.5), 1, 1, color=renkler[e, o], alpha=0.55, lw=0))
        hucre = df.groupby([o_k, e_k])["No"].apply(list)
        for (o, e), ids in hucre.items():
            for i, rid in enumerate(ids):
                ax.text(o - 0.38 + (i % 3) * 0.3, e + 0.3 - (i // 3) * 0.2, rid, fontsize=7.5, color=g.MUREKKEP, fontweight="bold")
        ax.set_xlim(0.5, 5.5); ax.set_ylim(0.5, 5.5); ax.set_xticks(range(1, 6)); ax.set_yticks(range(1, 6))
        ax.set_xlabel("Olasılık"); ax.set_ylabel("Etki"); ax.set_title(baslik_, loc="left")
        ax.set_aspect("equal")
        for k in ("top", "right"):
            ax.spines[k].set_visible(False)
    ust = df.sort_values(["Puan", "Etki (1-5)"], ascending=False).head(5)
    alt = "En yüksek 5: " + "; ".join(f"{r.No} {r.Risk} ({r.Puan})" for r in ust.itertuples())
    g.baslik(fig, "Risk ısı haritası: azaltımla en çok küçülen riskler takvim ve fiyat riskleri", alt)
    fig.tight_layout(rect=(0, 0, 1, 0.9)); g.kaynak_notu(fig, "Kaynak: kendi değerlendirmem (puanlar uzman yargısı, TAHMİN); sayısal etkiler model çıktılarından.")
    fig.savefig(yol); plt.close(fig)


def main():
    CIKTI.mkdir(parents=True, exist_ok=True)
    df = tablo()
    df.to_csv(CIKTI / "risk_kaydi.csv", encoding="utf-8-sig", index=False)
    (KOK / "results").mkdir(exist_ok=True)
    df.to_csv(KOK / "results" / "risk_kaydi.csv", encoding="utf-8-sig", index=False)
    excel(df, CIKTI / "risk_kaydi.xlsx")
    isi_haritasi(df, CIKTI / "risk_isi_haritasi.png")
    ust = df.sort_values(["Puan", "Etki (1-5)"], ascending=False).head(5)
    print(ust[["No", "Kategori", "Risk", "Puan", "Kalıntı puan"]].to_string(index=False))


if __name__ == "__main__":
    main()
