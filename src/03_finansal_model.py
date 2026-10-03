"""
Aşama 3 — Finansal model: referans konfigürasyonlar, senaryolar, duyarlılık, marjinal GES değeri

Konfigürasyonlar (50 MW şebeke bağlantısı, bataryasız):
    (a) RES 50 MW   (b) RES 50 + GES 25 MW   (c) RES 50 + GES 50 MW   (d) GES 50 MWac

Çıktılar: outputs/03_finans/*.png, *.csv, nakit_akisi_b.xlsx
Çalıştırma: .venv/bin/python src/03_finansal_model.py
"""

import copy
import sys
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from src import gorsel as g
from src.ayarlar import CIKTILAR, config_oku, deger
from src.finans.borc import borc_boyutlandir
from src.finans.fiyat_senaryosu import fiyat_senaryosu_uret
from src.finans.model import Varsayimlar, p90_carpani, proje_degerle
from src.finans.veri import uretim_profili

CIKTI = CIKTILAR / "03_finans"
KONFIGLER = {
    "(a) RES 50": dict(res_mw=50, ges_mw=0, sebeke_mw=50),
    "(b) RES 50 + GES 25": dict(res_mw=50, ges_mw=25, sebeke_mw=50),
    "(c) RES 50 + GES 50": dict(res_mw=50, ges_mw=50, sebeke_mw=50),
    "(d) GES 50": dict(res_mw=0, ges_mw=50, sebeke_mw=50),
}
SEVIYELER = ["dusuk", "baz", "yuksek"]
VARYANTLAR = ["durur", "devam"]
SEVIYE_AD = {"dusuk": "Düşük", "baz": "Baz", "yuksek": "Yüksek"}
KAYNAK = ("Kaynak: kendi modelim (Aşama 1-3). Girdiler: EPİAŞ PTF, PVGIS, ERA5 (kalibre), "
          "IRENA 2025, BNEF 2025, Lazard 2025, EPDK/TEİAŞ 2026 tarifeleri. Reel 2026 USD.")
KONFIG_RENK = {"(a) RES 50": g.RENK_RES, "(b) RES 50 + GES 25": g.RENK_HIBRIT,
               "(c) RES 50 + GES 50": "#0d7a55", "(d) GES 50": g.RENK_GES}


def degerle(k, v, f, u, cfg):
    """NPV/IRR/LCOE + P90 ve P50'ye göre borç kapasitesi + öz kaynak IRR (P90 borcuyla, P50 nakit akışı)."""
    s50 = proje_degerle(k, v, f, u)
    r = p90_carpani(k, u, v)
    s90 = proje_degerle(k, v, f, u, gelir_carpani=r)
    b90 = borc_boyutlandir(s90, cfg, nakit=s50)
    b50 = borc_boyutlandir(s50, cfg)
    o = s50.ozet
    return s50, {
        "NPV [M USD]": s50.npv / 1e6, "Proje IRR (reel, vergi sonrası) [%]": 100 * s50.irr,
        "Öz kaynak IRR reel [%]": 100 * b90.ozkaynak_irr_reel, "Öz kaynak IRR nominal [%]": 100 * b90.ozkaynak_irr_nominal,
        "LCOE GES [USD/MWh]": s50.lcoe["ges"], "LCOE RES [USD/MWh]": s50.lcoe["res"],
        "LCOE hibrit (satılan) [USD/MWh]": s50.lcoe["hibrit_satilan"],
        "Yakalanan fiyat [USD/MWh]": o["Yakalanan fiyat [USD/MWh]"], "Ort. PTF [USD/MWh]": o["Ort. PTF [USD/MWh]"],
        "GES yakalanan [USD/MWh]": o["GES yakalanan [USD/MWh]"], "RES yakalanan [USD/MWh]": o["RES yakalanan [USD/MWh]"],
        "Kesinti [%]": o["Kesinti oranı [%]"], "Satılan enerji yıl 1 [GWh]": o["Satılan enerji, yıl 1 [GWh]"],
        "CAPEX [M USD]": o["CAPEX [M USD]"], "P90/P50": r,
        "Borç kapasitesi P90 [M USD]": b90.borc_reel / 1e6, "Borç oranı P90 [%]": 100 * b90.borc_orani,
        "Borç kapasitesi P50 [M USD]": b50.borc_reel / 1e6, "Borç oranı P50 [%]": 100 * b50.borc_orani,
    }


# ---------------------------------------------------------------------------
def main():
    cfg = config_oku()
    CIKTI.mkdir(parents=True, exist_ok=True)
    u = uretim_profili()
    v0 = Varsayimlar.configten(cfg)

    senaryolar = {(s, kv): fiyat_senaryosu_uret(cfg, u, s, kv) for s in SEVIYELER for kv in VARYANTLAR}
    baz = senaryolar[("baz", "devam")]

    # --- 1) Konfigürasyonlar × senaryolar
    satirlar, sonuclar_baz = [], {}
    for ad, k in KONFIGLER.items():
        for (s, kv), f in senaryolar.items():
            sonuc, metrik = degerle(k, v0, f, u, cfg)
            satirlar.append({"Konfigürasyon": ad, "Fiyat seviyesi": SEVIYE_AD[s], "Kanibalizasyon": kv, **metrik})
            if (s, kv) == ("baz", "devam"):
                sonuclar_baz[ad] = sonuc
    tablo = pd.DataFrame(satirlar)
    tablo.to_csv(CIKTI / "konfigurasyon_senaryolar.csv", encoding="utf-8-sig", index=False, float_format="%.3f")
    ana = tablo[(tablo["Fiyat seviyesi"] == "Baz") & (tablo["Kanibalizasyon"] == "devam")].set_index("Konfigürasyon")
    ana.drop(columns=["Fiyat seviyesi", "Kanibalizasyon"]).to_csv(CIKTI / "konfigurasyon_baz.csv", encoding="utf-8-sig", float_format="%.3f")
    npv_tablo = tablo.pivot_table(index="Konfigürasyon", columns=["Kanibalizasyon", "Fiyat seviyesi"], values="NPV [M USD]")
    npv_tablo = npv_tablo.reindex(columns=pd.MultiIndex.from_product([VARYANTLAR, [SEVIYE_AD[s] for s in SEVIYELER]]))

    # --- 2) Tornado (konfigürasyon b, baz / devam)
    kb = KONFIGLER["(b) RES 50 + GES 25"]
    npv_b = sonuclar_baz["(b) RES 50 + GES 25"].npv / 1e6

    def npv_ile(v=None, f=None):
        return proje_degerle(kb, v or v0, f or baz, u).npv / 1e6

    kan = deger(cfg["fiyat_senaryolari"]["kanibalizasyon"])
    cfg_derin = copy.deepcopy(cfg)
    cfg_derin["fiyat_senaryolari"]["kanibalizasyon"]["value"] = {**kan, "taban_oran": 0.50}
    tornado = pd.DataFrame([
        {"Değişken": "CAPEX (±%20)", "Düşük durum": "+%20", "Yüksek durum": "−%20",
         "NPV düşük": npv_ile(v=replace(v0, capex_carpani=1.2)), "NPV yüksek": npv_ile(v=replace(v0, capex_carpani=0.8))},
        {"Değişken": "Fiyat seviyesi (45 / 72 USD/MWh)", "Düşük durum": "45", "Yüksek durum": "72",
         "NPV düşük": npv_ile(f=senaryolar[("dusuk", "devam")]), "NPV yüksek": npv_ile(f=senaryolar[("yuksek", "devam")])},
        {"Değişken": "Kanibalizasyon (taban 0,50 / durur)", "Düşük durum": "taban 0,50", "Yüksek durum": "durur",
         "NPV düşük": npv_ile(f=fiyat_senaryosu_uret(cfg_derin, u, "baz", "devam")), "NPV yüksek": npv_ile(f=senaryolar[("baz", "durur")])},
        {"Değişken": "WACC (%8,5 / %5,5)", "Düşük durum": "%8,5", "Yüksek durum": "%5,5",
         "NPV düşük": npv_ile(v=replace(v0, wacc=0.085)), "NPV yüksek": npv_ile(v=replace(v0, wacc=0.055))},
        {"Değişken": "RES kapasite faktörü (±%10)", "Düşük durum": "−%10", "Yüksek durum": "+%10",
         "NPV düşük": npv_ile(v=replace(v0, res_uretim_carpani=0.9)), "NPV yüksek": npv_ile(v=replace(v0, res_uretim_carpani=1.1))},
        {"Değişken": "Dengesizlik maliyeti (RES %8 / %2)", "Düşük durum": "RES %8, GES %4", "Yüksek durum": "RES %2, GES %1",
         "NPV düşük": npv_ile(v=replace(v0, dengesizlik={"res": 0.08, "ges": 0.04})),
         "NPV yüksek": npv_ile(v=replace(v0, dengesizlik={"res": 0.02, "ges": 0.01}))},
    ])
    tornado["Aralık"] = (tornado["NPV yüksek"] - tornado["NPV düşük"]).abs()
    tornado = tornado.sort_values("Aralık", ascending=False).reset_index(drop=True)
    tornado.to_csv(CIKTI / "tornado_b.csv", encoding="utf-8-sig", index=False, float_format="%.3f")

    # --- 3) Ek GES'in marjinal değeri (RES 50, şebeke 50; GES 0–75 MW)
    ges_mw = np.arange(0, 80, 5)
    marj = []
    for kv in VARYANTLAR:
        for gmw in ges_mw:
            s = proje_degerle(dict(res_mw=50, ges_mw=gmw, sebeke_mw=50), v0, senaryolar[("baz", kv)], u)
            y = s.yillik
            marj.append({"Kanibalizasyon": kv, "GES [MW]": gmw, "NPV [M USD]": s.npv / 1e6,
                         "GES üretim y1 [MWh]": y["GES üretim [MWh]"].iloc[0], "Satılan y1 [MWh]": y["Satılan [MWh]"].iloc[0],
                         "Gelir ömür toplamı [M USD]": y["Toplam gelir [USD]"].sum() / 1e6,
                         "Kesinti [%]": s.ozet["Kesinti oranı [%]"]})
    marj = pd.DataFrame(marj)
    # Bağımsız GES referansı: kendi bağlantısı olan GES'in MW başına NPV'si (konfig d / 50)
    ges_tek = {kv: proje_degerle(KONFIGLER["(d) GES 50"], v0, senaryolar[("baz", kv)], u).npv / 1e6 / 50 for kv in VARYANTLAR}
    adim = []
    for kv in VARYANTLAR:
        m = marj[marj["Kanibalizasyon"] == kv].set_index("GES [MW]")
        for g0, g1, ad in [(0, 25, "(a) → (b)"), (25, 50, "(b) → (c)")]:
            d_ges = m.loc[g1, "GES üretim y1 [MWh]"] - m.loc[g0, "GES üretim y1 [MWh]"]
            d_sat = m.loc[g1, "Satılan y1 [MWh]"] - m.loc[g0, "Satılan y1 [MWh]"]
            adim.append({"Kanibalizasyon": kv, "Adım": ad,
                         "ΔNPV / MW GES [M USD/MW]": (m.loc[g1, "NPV [M USD]"] - m.loc[g0, "NPV [M USD]"]) / (g1 - g0),
                         "Eklenen GES enerjisinin kesilen payı [%]": 100 * (1 - d_sat / d_ges),
                         "Bağımsız GES NPV / MW [M USD/MW]": ges_tek[kv]})
    adim = pd.DataFrame(adim)
    marj.to_csv(CIKTI / "marjinal_ges_tarama.csv", encoding="utf-8-sig", index=False, float_format="%.3f")
    adim.to_csv(CIKTI / "marjinal_ges_adimlar.csv", encoding="utf-8-sig", index=False, float_format="%.4f")

    # --- 4) Grafikler
    grafik_konfig(npv_tablo, ana)
    grafik_selale(sonuclar_baz["(b) RES 50 + GES 25"], v0)
    grafik_tornado(tornado, npv_b)
    grafik_lcoe(ana)
    grafik_fiyat(senaryolar, u)
    grafik_marjinal(marj, ges_tek)

    # --- 5) Excel (konfigürasyon b)
    excel_b(sonuclar_baz["(b) RES 50 + GES 25"], ana.loc["(b) RES 50 + GES 25"], v0, cfg, u, baz)

    pd.set_option("display.width", 250); pd.set_option("display.float_format", lambda x: f"{x:.2f}")
    print("=== Baz fiyat, kanibalizasyon devam ===")
    print(ana.drop(columns=["Fiyat seviyesi", "Kanibalizasyon"]).T.to_string())
    print("\n=== NPV [M USD] — senaryolar ==="); print(npv_tablo.to_string())
    print("\n=== Tornado (b), baz NPV = %.2f M USD ===" % npv_b); print(tornado.to_string(index=False))
    print("\n=== Marjinal GES ==="); print(adim.to_string(index=False))
    print("\nGüneş yakalanma oranı hedefi (devam):", np.round(baz.gunes_hedef_oran[[0, 5, 10, 29]], 3),
          "| başlangıç:", round(baz.bilgi["baslangic_gunes_orani"], 3))


# ---------------------------------------------------------------------------
def grafik_konfig(npv_tablo, ana):
    fig, axs = plt.subplots(1, 2, figsize=(13, 4.8), gridspec_kw={"width_ratios": [1.4, 1]})
    ax = axs[0]
    x = np.arange(len(npv_tablo)); w = 0.25
    tonlar = {"Düşük": "#9ec5ef", "Baz": g.RENK_RES, "Yüksek": "#1d4f91"}
    for i, s in enumerate(["Düşük", "Baz", "Yüksek"]):
        val = npv_tablo[("devam", s)].to_numpy()
        c = ax.bar(x + (i - 1) * (w + 0.02), val, w, color=tonlar[s], label=f"{s} fiyat")
        ax.bar_label(c, fmt="%.0f", fontsize=7, color=g.IKINCIL, padding=8)
        ax.scatter(x + (i - 1) * (w + 0.02), npv_tablo[("durur", s)], marker="_", s=180, color=g.MUREKKEP, zorder=3,
                   label="Kanibalizasyon durursa" if i == 0 else None)
    ax.axhline(0, color=g.MUREKKEP, lw=0.8)
    ax.set_xticks(x, npv_tablo.index)
    ax.set_ylabel("NPV [M USD, reel, %7 WACC]")
    ax.set_title("NPV — kanibalizasyon devam (çubuk) ve durur (çizgi)", loc="left")
    ax.legend(fontsize=8, ncols=2, loc="upper right")
    g.eksen_sade(ax)
    ax = axs[1]
    for ad, r in ana.iterrows():
        ax.scatter(r["LCOE hibrit (satılan) [USD/MWh]"], r["Yakalanan fiyat [USD/MWh]"], s=90, color=KONFIG_RENK[ad],
                   edgecolor="white", linewidth=1.5, zorder=3)
        ofset = {"(a) RES 50": (8, 4), "(b) RES 50 + GES 25": (-8, 8), "(c) RES 50 + GES 50": (8, -10), "(d) GES 50": (8, -3)}[ad]
        ax.annotate(ad, (r["LCOE hibrit (satılan) [USD/MWh]"], r["Yakalanan fiyat [USD/MWh]"]), xytext=ofset,
                    textcoords="offset points", fontsize=8, color=g.MUREKKEP, ha="right" if ofset[0] < 0 else "left")
    lim = [30, 62]
    ax.plot(lim, lim, color=g.SOLUK, lw=1, ls="--")
    ax.text(57, 59, "yakalanan = LCOE", fontsize=7.5, color=g.IKINCIL, ha="right")
    ax.set_xlim(*lim); ax.set_ylim(*lim)
    ax.set_xlabel("LCOE, satılan enerji [USD/MWh]"); ax.set_ylabel("Yakalanan fiyat [USD/MWh]")
    ax.set_title("Baz senaryo: çizginin üstü = maliyetini karşılıyor", loc="left")
    g.eksen_sade(ax, izgara="both")
    g.baslik(fig, "Hiçbir bataryasız konfigürasyon baz senaryoda %7 reel getiriyi aşamıyor",
             "50 MW şebeke bağlantısı; NPV üç fiyat seviyesi ve iki kanibalizasyon varyantı için. LCOE ve yakalanan fiyat ömür boyu ortalama.")
    fig.tight_layout(rect=(0, 0, 1, 0.88)); g.kaynak_notu(fig, KAYNAK)
    fig.savefig(CIKTI / "konfigurasyon_karsilastirma.png"); plt.close(fig)


def grafik_selale(s, v):
    """Ömür boyu bugünkü değer şelalesi (konfigürasyon b): gelirden NPV'ye."""
    y = s.yillik
    o = s.ozet
    nb = len(o["capex_akisi"]) - len(y)
    df = (1 + v.wacc) ** -(np.arange(nb, nb + len(y)) + 1)
    pv = lambda c: float((y[c].to_numpy() * df).sum()) / 1e6
    kalemler = [
        ("GES geliri", pv("GES geliri [USD]")), ("RES geliri", pv("RES geliri [USD]") - pv("Eşzamanlılık düzeltmesi [USD]")),
        ("Eşzamanlılık\ndüzeltmesi", pv("Eşzamanlılık düzeltmesi [USD]")), ("Dengesizlik", pv("Dengesizlik [USD]")),
        ("İşletme-bakım\n+ arazi", pv("İşletme-bakım + arazi [USD]")), ("Sigorta", pv("Sigorta [USD]")),
        ("TEİAŞ\nbedelleri", pv("TEİAŞ bedelleri [USD]")), ("Vergi", pv("Vergi [USD]")),
        ("Söküm", pv("Söküm [USD]")),
        ("CAPEX", -float((o["capex_akisi"] * (1 + v.wacc) ** -(np.arange(len(o["capex_akisi"])) + 1)).sum()) / 1e6),
    ]
    kalemler = [(ad, d) for ad, d in kalemler if abs(d) >= 0.05]
    fig, ax = plt.subplots(figsize=(11, 4.8))
    taban = 0.0
    for i, (ad, deg) in enumerate(kalemler):
        renk = g.RENK_GES if ad.startswith("GES") else g.RENK_RES if ad.startswith("RES") else ("#8a8984" if deg < 0 else g.RENK_HIBRIT)
        ax.bar(i, deg, 0.65, bottom=taban, color=renk)
        ax.text(i, taban + deg + (0.8 if deg >= 0 else -0.8), f"{deg:+.1f}", ha="center", va="bottom" if deg >= 0 else "top",
                fontsize=8, color=g.MUREKKEP)
        taban += deg
    ax.bar(len(kalemler), taban, 0.65, color=g.MUREKKEP)
    ax.text(len(kalemler), taban + (0.8 if taban >= 0 else -0.8), f"{taban:+.1f}", ha="center",
            va="bottom" if taban >= 0 else "top", fontsize=9, color=g.MUREKKEP, fontweight="bold")
    ax.axhline(0, color=g.MUREKKEP, lw=0.8)
    ax.set_xticks(range(len(kalemler) + 1), [k for k, _ in kalemler] + ["NPV"], fontsize=8)
    ax.set_ylabel("Bugünkü değer [M USD, %7 reel]")
    g.eksen_sade(ax)
    g.baslik(fig, "Konfigürasyon (b): gelirin bugünkü değeri ve NPV'ye köprü",
             "RES 50 MW + GES 25 MW, 50 MW bağlantı; baz fiyat, kanibalizasyon devam. Tüm kalemler FID'ye (2026) indirgenmiş.")
    fig.tight_layout(rect=(0, 0, 1, 0.9)); g.kaynak_notu(fig, KAYNAK)
    fig.savefig(CIKTI / "nakit_akisi_selale_b.png"); plt.close(fig)


def grafik_tornado(t, npv_b):
    fig, ax = plt.subplots(figsize=(10, 4.6))
    t = t.iloc[::-1].reset_index(drop=True)
    for i, r in t.iterrows():
        dus, yuk = r["NPV düşük"] - npv_b, r["NPV yüksek"] - npv_b
        ax.barh(i, dus, left=npv_b, color="#8a8984", height=0.6)
        ax.barh(i, yuk, left=npv_b, color=g.RENK_HIBRIT, height=0.6)
        ax.text(min(r["NPV düşük"], r["NPV yüksek"]) - 0.4, i, r["Düşük durum"] if r["NPV düşük"] < r["NPV yüksek"] else r["Yüksek durum"],
                ha="right", va="center", fontsize=7.5, color=g.IKINCIL)
        ax.text(max(r["NPV düşük"], r["NPV yüksek"]) + 0.4, i, r["Yüksek durum"] if r["NPV yüksek"] > r["NPV düşük"] else r["Düşük durum"],
                ha="left", va="center", fontsize=7.5, color=g.IKINCIL)
    ax.axvline(npv_b, color=g.MUREKKEP, lw=1)
    ax.axvline(0, color=g.MUREKKEP, lw=0.8, ls="--")
    ax.set_yticks(range(len(t)), t["Değişken"])
    ax.set_xlabel("NPV [M USD, reel]")
    lo, hi = t[["NPV düşük", "NPV yüksek"]].min().min(), t[["NPV düşük", "NPV yüksek"]].max().max()
    ax.set_xlim(lo - 6, hi + 6)
    g.eksen_sade(ax, izgara="x")
    g.baslik(fig, "Konfigürasyon (b) NPV duyarlılığı: fiyat seviyesi ve CAPEX belirleyici",
             f"Baz NPV {npv_b:.1f} M USD (dikey çizgi). Gri: olumsuz durum, yeşil: olumlu durum; tek seferde bir değişken.")
    fig.tight_layout(rect=(0, 0, 1, 0.9)); g.kaynak_notu(fig, KAYNAK)
    fig.savefig(CIKTI / "tornado_b.png"); plt.close(fig)


def grafik_lcoe(ana):
    fig, ax = plt.subplots(figsize=(10, 4.6))
    kalemler = []
    for ad, r in ana.iterrows():
        if np.isfinite(r["LCOE RES [USD/MWh]"]):
            kalemler.append((f"{ad[:3]}\nRES", r["LCOE RES [USD/MWh]"], g.RENK_RES))
        if np.isfinite(r["LCOE GES [USD/MWh]"]):
            kalemler.append((f"{ad[:3]}\nGES", r["LCOE GES [USD/MWh]"], g.RENK_GES))
        kalemler.append((f"{ad[:3]}\nsatılan", r["LCOE hibrit (satılan) [USD/MWh]"], g.MUREKKEP))
    for i, (ad, deg, renk) in enumerate(kalemler):
        ax.bar(i, deg, 0.65, color=renk)
        ax.text(i, deg + 0.6, f"{deg:.0f}", ha="center", fontsize=8, color=g.MUREKKEP)
    for y, ad in [(33, "IRENA 2025 küresel RES"), (44, "IRENA 2025 küresel GES"), (52, "IRENA 2025 Almanya RES"), (65, "IRENA 2025 Almanya GES")]:
        ax.hlines(y, -0.6, len(kalemler) - 0.4, color=g.SOLUK, lw=0.8, ls="--")
        ax.text(len(kalemler) - 0.3, y, ad, ha="left", va="center", fontsize=7.5, color=g.IKINCIL)
    ax.set_xlim(-0.6, len(kalemler) + 3.2)
    ax.set_xticks(range(len(kalemler)), [k for k, _, _ in kalemler], fontsize=8)
    ax.set_ylabel("LCOE [USD/MWh, reel, vergi öncesi]")
    ax.set_ylim(0, 72)
    g.eksen_sade(ax)
    g.baslik(fig, "LCOE: TEİAŞ bedelleri, arazi ve bağlantı maliyeti küresel ortalamanın üstüne itiyor",
             "(a) RES 50, (b) RES 50 + GES 25, (c) RES 50 + GES 50, (d) GES 50. Teknoloji LCOE'si kendi üretimi (kesinti öncesi); 'satılan' = toplam maliyet / satılan enerji. %7 reel WACC.")
    fig.tight_layout(rect=(0, 0, 1, 0.9)); g.kaynak_notu(fig, KAYNAK + " Referans: IRENA RPGC 2025.")
    fig.savefig(CIKTI / "lcoe_karsilastirma.png"); plt.close(fig)


def grafik_fiyat(senaryolar, u):
    fig, axs = plt.subplots(1, 3, figsize=(14, 4.4))
    ax = axs[0]
    for s, renk in zip(SEVIYELER, ["#9ec5ef", g.RENK_RES, "#1d4f91"]):
        f = senaryolar[(s, "devam")]
        ax.plot(f.yillar, f.fiyatlar.mean(axis=(1, 2)), color=renk, lw=2, label=f"{SEVIYE_AD[s]}: {f.seviye:.0f}")
    ax.set_ylim(0, 85); ax.set_ylabel("Yıllık ortalama PTF [USD/MWh, reel]")
    ax.set_title("Seviye senaryoları", loc="left"); ax.legend(fontsize=8)
    ax = axs[1]
    for kv, stil in [("durur", "--"), ("devam", "-")]:
        f = senaryolar[("baz", kv)]
        ax.plot(f.yillar, 100 * f.gunes_hedef_oran, color=g.RENK_GES, lw=2, ls=stil, label=f"Kanibalizasyon {kv}")
    ax.set_ylim(50, 90); ax.set_ylabel("Ezine GES yakalanma oranı [%]")
    ax.set_title("Kanibalizasyon varyantları", loc="left"); ax.legend(fontsize=8)
    ax = axs[2]
    saat = (np.arange(8760) + 3) % 24
    for kv, i, renk, stil, ad in [("devam", 0, g.SOLUK, "-", "2028"), ("devam", 5, g.MUREKKEP, "-", "2033+ (devam)")]:
        f = senaryolar[("baz", kv)]
        prof = pd.Series(f.fiyatlar[i].mean(axis=0)).groupby(saat).mean()
        ax.plot(prof.index, prof.values, color=renk, lw=2, ls=stil, label=ad)
    ax.set_xticks(range(0, 24, 3)); ax.set_xlabel("Saat (TSİ)"); ax.set_ylabel("Ortalama fiyat [USD/MWh]")
    ax.set_title("Baz senaryo gün içi profil", loc="left"); ax.legend(fontsize=8)
    for a in axs:
        g.eksen_sade(a, izgara="both")
    g.baslik(fig, "Fiyat senaryoları: seviye × kanibalizasyon",
             "Saatlik şekil 2025 ve Eki 2025–Eyl 2026 PTF'sinden; kanibalizasyon öğle çukurunu derinleştirir, yıllık ortalama sabit kalır.")
    fig.tight_layout(rect=(0, 0, 1, 0.88)); g.kaynak_notu(fig, KAYNAK)
    fig.savefig(CIKTI / "fiyat_senaryolari.png"); plt.close(fig)


def grafik_marjinal(marj, ges_tek):
    fig, axs = plt.subplots(1, 2, figsize=(13, 4.6))
    ax = axs[0]
    for kv, stil in [("durur", "--"), ("devam", "-")]:
        m = marj[marj["Kanibalizasyon"] == kv]
        ax.plot(m["GES [MW]"], m["NPV [M USD]"], color=g.RENK_HIBRIT, lw=2, ls=stil, marker="o", ms=4,
                markeredgecolor="white", label=f"Kanibalizasyon {kv}")
    for x, ad in [(0, "(a)"), (25, "(b)"), (50, "(c)")]:
        ax.axvline(x, color=g.IZGARA, lw=1)
        ax.text(x + 0.8, ax.get_ylim()[1], ad, va="top", fontsize=8, color=g.IKINCIL)
    ax.set_xlabel("Eklenen GES [MWac] (RES 50 MW, şebeke 50 MW)"); ax.set_ylabel("NPV [M USD]")
    ax.set_title("GES eklendikçe NPV düşüyor", loc="left"); ax.legend(fontsize=8)
    ax = axs[1]
    m = marj[marj["Kanibalizasyon"] == "devam"].set_index("GES [MW]")
    marjinal = m["NPV [M USD]"].diff() / 5
    ax.bar(marjinal.index[1:] - 2.5, marjinal.values[1:], 4, color=g.RENK_GES)
    ax.axhline(ges_tek["devam"], color=g.MUREKKEP, lw=1.2, ls="--")
    ax.text(-1.5, ges_tek["devam"] + 0.008, "bağımsız GES (kendi bağlantısı)", ha="left", va="bottom", fontsize=7.5, color=g.MUREKKEP)
    ax.axhline(0, color=g.MUREKKEP, lw=0.8)
    isaret = [x for x in m.index if x % 15 == 0]
    ax.set_xticks(isaret, [f"{x}\nkesinti %{m.loc[x, 'Kesinti [%]']:.0f}" for x in isaret], fontsize=8)
    ax.set_xlabel("GES [MWac] ve toplam kesinti oranı"); ax.set_ylabel("Her ek MW GES'in NPV katkısı [M USD/MW]")
    ax.set_title("Marjinal değer: kesinti arttıkça her ek MW daha az değerli", loc="left")
    for a in axs:
        g.eksen_sade(a)
    g.baslik(fig, "Ek GES'in marjinal değeri: kesinti + kanibalizasyon",
             "Baz fiyat. Sağ: 5 MW'lık adımların MW başına NPV değişimi; kesikli çizgi: kendi bağlantısı olan bağımsız GES'in MW başına NPV'si.")
    fig.tight_layout(rect=(0, 0, 1, 0.88)); g.kaynak_notu(fig, KAYNAK)
    fig.savefig(CIKTI / "marjinal_ges.png"); plt.close(fig)


# ---------------------------------------------------------------------------
def excel_b(s, metrik, v, cfg, u, f):
    """Konfigürasyon (b) yıllık nakit akışı — finansçıların inceleyebileceği biçimde."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    k = KONFIGLER["(b) RES 50 + GES 25"]
    r = p90_carpani(k, u, v)
    b90 = borc_boyutlandir(proje_degerle(k, v, f, u, gelir_carpani=r), cfg, nakit=s)
    o = s.ozet
    nb = v.insaat_suresi
    yil0 = deger(cfg["proje"]["fid_yili"])

    # Zaman çizelgesi tablosu (inşaat + işletme) — USD bin
    y = s.yillik.copy()
    ins = pd.DataFrame({"İşletme yılı": [f"İnşaat {i + 1}" for i in range(nb)], "Takvim yılı": [yil0 + i for i in range(nb)]})
    tablo = pd.concat([ins, y.astype({"İşletme yılı": str})], ignore_index=True)
    tablo["CAPEX [USD]"] = -o["capex_akisi"]
    tablo["Serbest nakit akışı [USD]"] = o["nakit_akisi"]
    tablo["İskonto çarpanı"] = (1 + v.wacc) ** -(np.arange(len(tablo)) + 1)
    tablo["İndirgenmiş SNA [USD]"] = tablo["Serbest nakit akışı [USD]"] * tablo["İskonto çarpanı"]
    servis = np.concatenate([np.zeros(nb), b90.servis_nominal])
    tablo["Borç servisi, nominal [USD]"] = -servis
    tablo["Öz kaynak nakit akışı, reel [USD]"] = b90.ozkaynak_akisi
    tablo = tablo.fillna(0)
    para = [c for c in tablo.columns if "[USD]" in c]
    tablo[para] = tablo[para] / 1000
    tablo = tablo.rename(columns={c: c.replace("[USD]", "[bin USD]") for c in para})

    ozet = pd.DataFrame({
        "Gösterge": ["Konfigürasyon", "Fiyat senaryosu", "NPV (reel, %7 WACC) [M USD]", "Proje IRR (reel, vergi sonrası) [%]",
                     "Öz kaynak IRR reel / nominal [%]", "CAPEX [M USD]", "LCOE GES / RES / satılan [USD/MWh]",
                     "Yakalanan fiyat (ömür ort.) [USD/MWh]", "Kesinti [%]", "P90 / P50 satılan enerji",
                     "Borç kapasitesi (P90, DSCR 1,30, 15 yıl) [M USD]", "Borç / CAPEX [%]"],
        "Değer": ["RES 50 MW + GES 25 MWac, 50 MW bağlantı", f.ad, round(s.npv / 1e6, 2), round(100 * s.irr, 2),
                  f"{100 * b90.ozkaynak_irr_reel:.1f} / {100 * b90.ozkaynak_irr_nominal:.1f}", round(o["CAPEX [M USD]"], 2),
                  f"{s.lcoe['ges']:.1f} / {s.lcoe['res']:.1f} / {s.lcoe['hibrit_satilan']:.1f}",
                  round(o["Yakalanan fiyat [USD/MWh]"], 1), round(o["Kesinti oranı [%]"], 2), round(r, 3),
                  round(b90.borc_reel / 1e6, 2), round(100 * b90.borc_orani, 1)],
    })
    varsayim = pd.DataFrame([
        ("WACC (reel, USD)", v.wacc), ("Kurumlar vergisi", v.vergi), ("CAPEX GES [USD/MWac]", v.capex_ges),
        ("CAPEX RES [USD/MW]", v.capex_res), ("CAPEX şebeke [USD/MW]", v.capex_sebeke),
        ("OPEX GES (+arazi) [USD/MW/yıl]", v.opex_ges + v.arazi_ges), ("OPEX RES (+arazi) [USD/MW/yıl]", v.opex_res + v.arazi_res),
        ("Sigorta [CAPEX %/yıl]", v.sigorta_oran), ("TEİAŞ kapasite [USD/MW/yıl]", v.teias_kapasite),
        ("TEİAŞ enerji [USD/MWh]", v.teias_enerji), ("Dengesizlik RES / GES [gelir %]", f"{v.dengesizlik['res']} / {v.dengesizlik['ges']}"),
        ("Degradasyon GES / RES [%/yıl]", f"{v.degr_ges} / {v.degr_res}"), ("Amortisman [yıl, doğrusal]", 10),
        ("Fiyat seviyesi [USD/MWh, reel]", f.seviye), ("Kanibalizasyon", f.bilgi["varyant"]),
    ], columns=["Varsayım", "Değer"])

    yol = CIKTI / "nakit_akisi_b.xlsx"
    with pd.ExcelWriter(yol, engine="openpyxl") as w:
        ozet.to_excel(w, sheet_name="Özet", index=False)
        tablo.to_excel(w, sheet_name="Nakit akışı", index=False)
        varsayim.to_excel(w, sheet_name="Varsayımlar", index=False)
        baslik = Font(bold=True, color="FFFFFF"); dolgu = PatternFill("solid", fgColor="1D4F91")
        ince = Side(style="thin", color="D0D0D0")
        for ad, ws in w.sheets.items():
            for hucre in ws[1]:
                hucre.font, hucre.fill = baslik, dolgu
                hucre.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.row_dimensions[1].height = 42
            for i, kolon in enumerate(ws.columns, start=1):
                genislik = max(len(str(c.value)) if c.value is not None else 0 for c in kolon)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, genislik * 0.9), 48)
            ws.freeze_panes = "C2" if ad == "Nakit akışı" else "A2"
        ws = w.sheets["Nakit akışı"]
        for satir in ws.iter_rows(min_row=2):
            for hucre in satir:
                hucre.border = Border(bottom=ince)
                if isinstance(hucre.value, float):
                    hucre.number_format = "0.0000" if "İskonto" in str(ws.cell(1, hucre.column).value) else "#,##0;[Red]-#,##0"
        ws["A" + str(ws.max_row + 2)] = "Birimler: bin USD, reel 2026 (borç servisi nominal USD). Nakit akışları yıl sonu; t=0 FID yılı başı."


if __name__ == "__main__":
    main()
